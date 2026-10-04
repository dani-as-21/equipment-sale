"""Unified intake for equipment documents, offers, service quotes, and appraisals.

Photo capture stays on its own path. This module does not create inventory rows,
does not copy one component's facts onto its siblings, and does not let an offer
set or move the valuation range.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from pathlib import Path

from app.logic import (
    DOCUMENT_SUFFIXES,
    IMAGE_SUFFIXES,
    _bounded_keys,
    _read_upload_text,
    equipment_for_match,
    parse_json,
    run_ocr,
    upsert_review,
)
from app.machines import membership_for_tags
from app.matching import match_inventory

TYPE_LABELS = {
    "equipment": "מידע על ציוד",
    "offer": "הצעת רכש",
    "service_quote": "הצעת מחיר לשירות",
    "appraisal": "ראיית שווי",
    "mixed": "מסמך מעורב",
    "unclear": "סוג לא ברור",
}
STATUS_LABELS = {
    "filed": "עובד ושויך",
    "awaiting_match": "ממתין לשיוך פריט",
    "awaiting_clarification": "ממתין להבהרה",
    "offer_awaiting_valuation": "הצעה ממתינה לשווי",
    "offer_ready": "הצעה מוכנה לבדיקה",
    "failed": "העיבוד נכשל",
}
NEXT_LABELS = {
    "clarify_terms": "להבהיר תנאים",
    "complete_asset_information": "להשלים מידע על הנכס",
    "research_value": "לחפש ראיות שווי",
    "specialist": "לקבל חוות דעת מומחה",
    "negotiate": "לשקול משא ומתן",
    "consider_acceptance": "לשקול קבלה",
}
PHRASES = {
    "offer": (
        "purchase offer",
        "offer to purchase",
        "הצעת רכש",
        "הצעה לרכישה",
        "מציעים לרכוש",
    ),
    "service_quote": (
        "quotation for dismantling",
        "quotation for transport",
        "quotation for repair",
        "quotation for inspection",
        "service quotation",
        "dismantling quotation",
        "הצעת מחיר לפירוק",
        "הצעת מחיר להובלה",
        "הצעת מחיר לתיקון",
        "הצעת מחיר לבדיקה",
    ),
    "appraisal": ("appraisal", "market value", "שומה", "שווי שוק"),
    "equipment": (
        "maintenance record",
        "specification",
        "nameplate",
        "condition report",
        "מפרט טכני",
        "דוח תחזוקה",
        "דוח מצב",
    ),
}
MONEY_RE = re.compile(
    r"(?:(USD|EUR|ILS|GBP|NIS|₪|\$|€)\s*([0-9]{1,3}(?:[ ,][0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?))"
    r"|([0-9]{1,3}(?:[ ,][0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)\s*(USD|EUR|ILS|GBP|NIS|₪|\$|€)",
    re.IGNORECASE,
)
DATE_RE = re.compile(r"(\d{4}-\d{2}-\d{2}|\d{1,2}[./]\d{1,2}[./]\d{2,4})")
CURRENCY = {"$": "USD", "€": "EUR", "₪": "ILS", "NIS": "ILS", "USD": "USD", "EUR": "EUR", "ILS": "ILS", "GBP": "GBP"}


def _digest(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()[:16]


def _money(raw: str) -> float | None:
    cleaned = raw.replace(" ", "").replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def _currency(token: str) -> str:
    return CURRENCY.get((token or "").upper(), CURRENCY.get(token or "", ""))


def _find_money(text: str) -> list[dict]:
    found = []
    for match in MONEY_RE.finditer(text or ""):
        if match.group(1):
            currency, amount = _currency(match.group(1)), _money(match.group(2))
        else:
            amount, currency = _money(match.group(3)), _currency(match.group(4))
        if amount is None or not currency:
            continue
        found.append({"amount": amount, "currency": currency, "text": match.group(0).strip()})
    return found


def _labeled(text: str, labels: tuple[str, ...]) -> str:
    for label in labels:
        match = re.search(rf"(?:{label})\s*[:\-]\s*(.+)", text or "", re.IGNORECASE)
        if match:
            return match.group(1).strip().split("\n")[0][:240]
    return ""


def _dates_after(text: str, labels: tuple[str, ...]) -> str:
    for label in labels:
        match = re.search(rf"(?:{label})\s*[:\-]\s*(\d{{4}}-\d{{2}}-\d{{2}}|\d{{1,2}}[./]\d{{1,2}}[./]\d{{2,4}})", text or "", re.IGNORECASE)
        if match:
            return match.group(1)
    return ""


def classify_text(name: str, text: str) -> dict:
    haystack = f"{name}\n{text}".casefold()
    scores = {key: sum(haystack.count(phrase.casefold()) for phrase in phrases) for key, phrases in PHRASES.items()}
    ranked = sorted((score, key) for key, score in scores.items() if score)
    if not ranked:
        return {
            "doc_type": "unclear",
            "confidence": "uncertain",
            "reason": "לא נמצא ניסוח שמזהה מידע ציוד, הצעת רכש, הצעת שירות או שומה.",
        }
    top, second = ranked[0], ranked[1] if len(ranked) > 1 else (0, "")
    if second[0] and second[0] >= top[0]:
        return {
            "doc_type": "mixed",
            "confidence": "uncertain",
            "reason": f"יש סימנים גם ל{TYPE_LABELS[top[1]]} וגם ל{TYPE_LABELS[second[1]]}.",
        }
    if top[0] < 1:
        return {"doc_type": "unclear", "confidence": "uncertain", "reason": "הסימנים חלשים מדי."}
    return {
        "doc_type": top[1],
        "confidence": "clear",
        "reason": f"המסמך משתמש בניסוח של {TYPE_LABELS[top[1]]}.",
    }


def extract_document(path: Path) -> dict:
    suffix = path.suffix.lower()
    parts: list[dict] = []
    note = ""
    if suffix in IMAGE_SUFFIXES or suffix == "":
        ocr = run_ocr(str(path))
        text = ocr.get("raw_text") or ""
        if text.strip():
            parts.append({"ref": "התמונה", "text": text.strip()})
        elif not ocr.get("available"):
            note = ocr.get("message") or "הטקסט בתמונה לא נקרא."
        else:
            note = "התמונה נשמרה, ולא נקרא ממנה טקסט."
        return {"text": text.strip(), "parts": parts, "note": note}
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            for index, page in enumerate(reader.pages[:30], start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    parts.append({"ref": f"עמוד {index}", "text": page_text[:8000]})
            if not parts:
                note = "ה-PDF נשמר, ולא חולץ ממנו טקסט. ייתכן שהוא סרוק."
        except Exception:
            note = "ה-PDF נשמר, והטקסט לא חולץ."
        return {"text": "\n".join(part["text"] for part in parts)[:30000], "parts": parts, "note": note}
    if suffix in {".xlsx", ".csv"}:
        if suffix == ".csv":
            lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
            for index, line in enumerate(lines[:200], start=1):
                if line.strip():
                    parts.append({"ref": f"שורה {index}", "text": line.strip()[:500]})
        else:
            try:
                from openpyxl import load_workbook

                book = load_workbook(path, read_only=True, data_only=True)
                try:
                    for sheet in book.worksheets:
                        for index, row in enumerate(sheet.iter_rows(max_row=200, max_col=12, values_only=True), start=1):
                            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                            if values:
                                parts.append({"ref": f"גיליון {sheet.title} שורה {index}", "text": " | ".join(values)[:500]})
                finally:
                    book.close()
            except Exception:
                note = "הגיליון נשמר, והשורות לא נקראו."
        if not parts and not note:
            note = "הגיליון נשמר, ולא נמצאו בו תאים עם טקסט."
        return {"text": "\n".join(part["text"] for part in parts)[:30000], "parts": parts, "note": note}
    if suffix in {".docx", ".txt", ".rtf"}:
        text, read_note = _read_upload_text(path)
        paragraphs = [line.strip() for line in (text or "").splitlines() if line.strip()]
        for index, line in enumerate(paragraphs[:200], start=1):
            parts.append({"ref": f"פסקה {index}", "text": line[:800]})
        return {"text": (text or "")[:30000], "parts": parts, "note": read_note}
    if suffix in DOCUMENT_SUFFIXES:
        return {"text": "", "parts": [], "note": "סוג המסמך נשמר, והטקסט שלו לא חולץ."}
    return {"text": "", "parts": [], "note": "סוג הקובץ לא נתמך."}


def _indexes(conn: sqlite3.Connection) -> tuple[list[dict], dict, dict]:
    equipment = equipment_for_match(conn)
    tags: dict[str, list[dict]] = {}
    serials: dict[str, list[dict]] = {}
    for row in equipment:
        if row.get("tag_norm"):
            tags.setdefault(row["tag_norm"], []).append(row)
        if row.get("serial_norm"):
            serials.setdefault(row["serial_norm"], []).append(row)
    return equipment, tags, serials


def _match_assets(conn: sqlite3.Connection, text: str) -> dict:
    equipment, tags, serials = _indexes(conn)
    exact_tags = _bounded_keys(text, tags)
    exact_serials = [token for token in _bounded_keys(text, serials) if token not in tags]
    confirmed = []
    candidates = []
    seen = set()
    for token in exact_tags + exact_serials:
        probed = match_inventory(equipment, tag=token if token in tags else "", serial=token if token in serials else "")
        if probed.get("mode") == "auto":
            for link in probed["links"]:
                if link["row_id"] in seen:
                    continue
                seen.add(link["row_id"])
                row = next(item for item in equipment if item["id"] == link["row_id"])
                confirmed.append({"row_id": row["id"], "tag_norm": row["tag_norm"], "tag": row["tag_original"], "reason": link["reason"]})
        else:
            for suggestion in (probed.get("suggestions") or [])[:8]:
                candidates.append({**suggestion, "tag_norm": token, "reason": suggestion.get("why") or "התג לא חד-משמעי."})
    return {"confirmed": confirmed, "candidates": candidates}


def _parts_for_tag(parts: list[dict], tag: str) -> list[dict]:
    needle = (tag or "").upper()
    return [part for part in parts if needle and needle in part["text"].upper()]


def _fact_from_part(part: dict) -> tuple[str, str]:
    text = part["text"]
    if re.search(r"condition|מצב", text, re.IGNORECASE):
        return "condition", text[:500]
    if re.search(r"maintenance|תחזוקה", text, re.IGNORECASE):
        return "maintenance", text[:500]
    if re.search(r"spec|מפרט|capacity|הספק", text, re.IGNORECASE):
        return "spec", text[:500]
    return "excerpt", text[:500]


def _remember_fact(conn, document_id: str, row_id: str, tag: str, key: str, text: str, source_ref: str, created_at: str) -> str:
    existing = conn.execute(
        """
        SELECT id, fact_text FROM intake_facts
        WHERE inventory_row_id = ? AND fact_key = ? AND fact_text = ?
        """,
        (row_id, key, text),
    ).fetchone()
    if existing:
        return ""
    conflict = conn.execute(
        """
        SELECT id, fact_text FROM intake_facts
        WHERE inventory_row_id = ? AND fact_key = ? AND fact_text != ?
        ORDER BY created_at DESC LIMIT 1
        """,
        (row_id, key, text),
    ).fetchone()
    fact_id = "IF-" + _digest(f"{document_id}|{row_id}|{key}|{text}")
    if conn.execute("SELECT id FROM intake_facts WHERE id = ?", (fact_id,)).fetchone():
        return ""
    conn.execute(
        """
        INSERT INTO intake_facts (
          id, document_id, inventory_row_id, tag_norm, fact_key, fact_text, source_ref, conflict_with, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (fact_id, document_id, row_id, tag, key, text, source_ref, conflict["id"] if conflict else "", created_at),
    )
    if not conflict:
        return ""
    upsert_review(
        conn,
        kind="conflict",
        queue="intake",
        dedupe_key=f"intake:{row_id}:{key}:conflict",
        question=f"ל{tag} יש שני ניסוחים שונים לשדה {key}. אף אחד לא דרס את השני.",
        priority="high",
        capture_id=None,
        row_ids=[row_id],
        payload={
            "document_id": document_id,
            "kept": conflict["fact_text"],
            "added": text,
            "resolve_by": "איזה ניסוח נכון, או שהשניים מתארים זמנים שונים. לא מוחקים את המקור.",
        },
        created_at=created_at,
    )
    return conflict["id"]


def _link(conn, document_id: str, row_id: str | None, tag: str, status: str, reason: str) -> None:
    link_id = "IL-" + _digest(f"{document_id}|{tag}|{row_id or ''}|{status}")
    if conn.execute("SELECT id FROM intake_links WHERE document_id = ? AND tag_norm = ? AND IFNULL(inventory_row_id, '') = ?", (document_id, tag, row_id or "")).fetchone():
        conn.execute(
            "UPDATE intake_links SET link_status = ?, reason = ? WHERE document_id = ? AND tag_norm = ? AND IFNULL(inventory_row_id, '') = ?",
            (status, reason, document_id, tag, row_id or ""),
        )
        return
    conn.execute(
        """
        INSERT INTO intake_links (id, document_id, inventory_row_id, tag_norm, link_status, reason)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (link_id, document_id, row_id, tag, status, reason),
    )


def _file_equipment(conn, document_id: str, parts: list[dict], matches: dict, created_at: str) -> list[str]:
    conflicts = []
    for item in matches["confirmed"]:
        _link(conn, document_id, item["row_id"], item["tag_norm"], "confirmed", item["reason"])
        for part in _parts_for_tag(parts, item["tag_norm"]) or []:
            key, text = _fact_from_part(part)
            conflict = _remember_fact(conn, document_id, item["row_id"], item["tag_norm"], key, text, part["ref"], created_at)
            if conflict:
                conflicts.append(item["tag_norm"])
    for item in matches["candidates"]:
        _link(conn, document_id, item.get("row_id"), item.get("tag_norm") or item.get("tag") or "", "candidate", item.get("reason") or "")
    return conflicts


def _price_basis(text: str, asset_count: int) -> str:
    haystack = (text or "").casefold()
    if any(phrase in haystack for phrase in ("indivisible", "lump sum", "package price", "חבילה", "לא ניתן לפיצול")):
        return "package"
    if any(phrase in haystack for phrase in ("per item", "each unit", "ליחידה")):
        return "per_item"
    if asset_count > 1 and any(phrase in haystack for phrase in ("total for", "for the listed", "עבור הפריטים")):
        return "several"
    if asset_count > 1:
        return "unknown"
    return "per_item" if asset_count == 1 and "offer price" in haystack else "unknown"


def _offer_fields(text: str, matches: dict) -> dict:
    labeled = _labeled(text, ("offer price", "purchase price", "total package price", "מחיר ההצעה", "סכום ההצעה"))
    labeled_money = _find_money(labeled) if labeled else []
    amounts = labeled_money or _find_money(text)
    amount = amounts[0]["amount"] if len(amounts) == 1 or labeled_money else None
    currency = amounts[0]["currency"] if amount is not None else ""
    vat = "unknown"
    haystack = text.casefold()
    if any(phrase in haystack for phrase in ("excluding vat", "excluded vat", "vat excluded", "plus vat", "בתוספת מע", "ללא מע")):
        vat = "excluded"
    elif any(phrase in haystack for phrase in ("including vat", "included vat", "vat included", "כולל מע")):
        vat = "included"
    exclusions = []
    for line in text.splitlines():
        if re.search(r"exclud|לא כולל|אינו כולל", line, re.IGNORECASE):
            exclusions.append(line.strip()[:240])
    included = matches["confirmed"]
    excluded_tags = []
    for line in exclusions:
        for item in included:
            if item["tag_norm"] and item["tag_norm"] in line.upper():
                excluded_tags.append(item["tag_norm"])
    assets = [item for item in included if item["tag_norm"] not in excluded_tags]
    basis = _price_basis(text, len(assets))
    missing = []
    buyer = _labeled(text, ("buyer", "הקונה", " purchaser"))
    if not buyer:
        missing.append("קונה")
    offer_date = _dates_after(text, ("offer date", "date", "תאריך"))
    expiry = _dates_after(text, ("expiry", "valid until", "expires", "בתוקף עד"))
    if not offer_date:
        missing.append("תאריך הצעה")
    if not expiry:
        missing.append("תוקף")
    if amount is None:
        missing.append("סכום")
    if not currency:
        missing.append("מטבע")
    if vat == "unknown":
        missing.append("טיפול במע״מ")
    payment = _labeled(text, ("payment", "תשלום"))
    if not payment:
        missing.append("תנאי תשלום")
    inspection = _labeled(text, ("inspection", "בדיקה"))
    contingencies = _labeled(text, ("contingenc", "תנאים מתלים", "subject to"))
    dismantling = _labeled(text, ("dismantling", "פירוק"))
    loading = _labeled(text, ("loading", "העמסה"))
    transport = _labeled(text, ("transport", "הובלה", "שינוע"))
    removal = _dates_after(text, ("removal", "פינוי"))
    for label, value in (("פירוק", dismantling), ("העמסה", loading), ("הובלה", transport)):
        if not value:
            missing.append(f"מי מטפל ב{label}")
    return {
        "buyer": buyer,
        "contacts": _labeled(text, ("contact", "איש קשר", "טלפון", "דוא\"ל", "email")),
        "offer_date": offer_date,
        "expiry": expiry,
        "price_basis": basis,
        "amount": amount,
        "currency": currency,
        "vat": vat,
        "payment_terms": payment,
        "inspection": inspection,
        "contingencies": contingencies,
        "dismantling": dismantling,
        "loading": loading,
        "transport": transport,
        "removal_date": removal,
        "exclusions": "\n".join(exclusions),
        "missing": missing,
        "assets": assets,
        "excluded_tags": excluded_tags,
        "candidate_tags": [item.get("tag_norm") or item.get("tag") for item in matches["candidates"]],
        "amounts_seen": amounts,
    }


def _latest_range(conn, subject_id: str) -> dict | None:
    row = conn.execute(
        """
        SELECT range_low, range_high, currency, range_label, confidence, confidence_why, created_at, assumptions
        FROM valuation_versions
        WHERE subject_id = ? AND scenario = 'normal_marketing' AND outdated = 0
        ORDER BY version_no DESC LIMIT 1
        """,
        (subject_id,),
    ).fetchone()
    return dict(row) if row else None


def _subject_for_rows(conn, row_ids: list[str]) -> str:
    if len(row_ids) == 1:
        from app.valuation import _id

        return _id("VAL-", row_ids[0])
    from app.valuation import _id

    return _id("VALG-", "|".join(sorted(row_ids)))


def _group_matches(conn, row_ids: list[str]) -> str:
    wanted = set(row_ids)
    for subject in conn.execute("SELECT id FROM valuation_subjects WHERE kind = 'group'"):
        members = {
            row["row_id"]
            for row in conn.execute(
                "SELECT row_id FROM valuation_members WHERE subject_id = ? AND role = 'included'",
                (subject["id"],),
            )
        }
        if members == wanted:
            return subject["id"]
    return ""


def _assess_offer(conn, fields: dict, created_at: str) -> dict:
    from app.valuation import ensure_row_subject

    row_ids = [item["row_id"] for item in fields["assets"] if item.get("row_id")]
    not_enough = "אין די מידע להערכה"
    differences = []
    evidence = []
    position = ""
    range_info = None
    subject_id = ""
    if not row_ids:
        conclusion = not_enough
        why = "הפריטים בהצעה לא זוהו בשורות המלאי, ולכן אין למה להשוות."
        action = "complete_asset_information"
        confidence = "low"
    elif fields["price_basis"] in {"package", "several", "unknown"} and len(row_ids) > 1:
        subject_id = _group_matches(conn, row_ids)
        for row_id in row_ids:
            ensure_row_subject(conn, row_id, created_at)
        if not subject_id:
            conclusion = not_enough
            why = "אין שווי לאותה חבילה. שווי הרכיבים לא סוכם ולא פוצל המחיר."
            action = "research_value"
            confidence = "low"
            differences.append("השווי הקיים, אם יש, הוא לרכיב ולא לחבילה שבהצעה.")
        else:
            range_info = _latest_range(conn, subject_id)
            conclusion, why, action, confidence, position = _place_amount(fields, range_info)
    else:
        if len(row_ids) > 1:
            conclusion = not_enough
            why = "לא ברור אם המחיר הוא לחבילה או לכל פריט, ולכן הוא לא פוצל ולא הושווה."
            action = "clarify_terms"
            confidence = "low"
        else:
            subject_id = ensure_row_subject(conn, row_ids[0], created_at)["id"]
            range_info = _latest_range(conn, subject_id)
            conclusion, why, action, confidence, position = _place_amount(fields, range_info)
    if fields["vat"] == "unknown":
        differences.append("טיפול המע״מ לא צוין, ולא נורמל.")
    if fields["currency"] and range_info and range_info.get("currency") and fields["currency"] != range_info["currency"]:
        differences.append("המטבעות שונים, ואין במסמך שער ותאריך המרה.")
        position = ""
        conclusion = not_enough
        why = "אי אפשר למקם את ההצעה מול הטווח בלי שער המרה שכתוב במסמך."
        action = "clarify_terms"
        confidence = "low"
    seller_costs = []
    unknown_costs = []
    for label, value in (("פירוק", fields["dismantling"]), ("העמסה", fields["loading"]), ("הובלה", fields["transport"])):
        if not value:
            unknown_costs.append(label)
        elif re.search(r"seller|המוכר|אנחנו|we will|on seller", value, re.IGNORECASE):
            seller_costs.append(label)
    net = None
    net_note = "עלויות המוכר לא ידועות, ולכן התמורה נטו לא חושבה ולא נרשמה כאפס."
    if unknown_costs:
        net_note = "חסר מי מטפל ב" + ", ".join(unknown_costs) + ". עלות לא ידועה לא נחשבה כאפס."
    elif seller_costs:
        net_note = "צוין שהמוכר נושא ב" + ", ".join(seller_costs) + ", ואין סכום לעלות הזו. התמורה נטו לא חושבה."
    if fields["exclusions"]:
        differences.append("ההצעה מציינת החרגות: " + fields["exclusions"][:240])
    if fields["excluded_tags"]:
        differences.append("ההצעה מוציאה רכיבים שהיו בזיהוי: " + ", ".join(fields["excluded_tags"]))
    if any(word in (fields["dismantling"] + fields["transport"]).casefold() for word in ("included", "כולל")):
        differences.append("ההצעה כוללת שירות שהשווי של הנכס בלבד לא כולל.")
    if position == "within":
        why += " מיקום בתוך הטווח אינו המלצה לקבל."
    if fields["missing"] and action == "consider_acceptance":
        action = "clarify_terms"
    recommendation = {
        "action": action,
        "label": NEXT_LABELS[action],
        "separated": "זו המלצה נפרדת. המערכת לא מאשרת, לא דוחה, לא יוצרת קשר ולא מתחייבת.",
    }
    return {
        "conclusion": conclusion,
        "position": position,
        "confidence": confidence,
        "confidence_why": why,
        "missing": fields["missing"],
        "differences": differences,
        "range": range_info,
        "subject_id": subject_id,
        "evidence_note": "ההצעה לא נכנסה לטווח ולא עיגנה אותו.",
        "market_evidence": evidence,
        "net_proceeds": net,
        "net_note": net_note,
        "recommendation": recommendation,
        "not_advice": position != "within" or True,
    }


def _place_amount(fields: dict, range_info: dict | None) -> tuple[str, str, str, str, str]:
    not_enough = "אין די מידע להערכה"
    if not range_info or range_info.get("range_low") is None:
        return not_enough, "אין טווח שנתמך בראיות לאותם נכסים.", "research_value", "low", ""
    if fields["amount"] is None or not fields["currency"]:
        return not_enough, "בלי סכום ומטבע אי אפשר להשוות לטווח.", "clarify_terms", "low", ""
    if fields["currency"] != (range_info.get("currency") or ""):
        return not_enough, "המטבע שונה מהטווח.", "clarify_terms", "low", ""
    low, high = range_info["range_low"], range_info["range_high"]
    if fields["amount"] < low:
        position, action = "below", "negotiate"
    elif fields["amount"] > high:
        position, action = "above", "consider_acceptance"
    else:
        position, action = "within", "consider_acceptance"
    labels = {"below": "מתחת לטווח", "within": "בתוך הטווח", "above": "מעל הטווח"}
    why = (
        f"ההצעה {labels[position]} של {low:g}–{high:g} {range_info.get('currency')} "
        f"מתאריך {str(range_info.get('created_at') or '')[:10]}. "
        f"רמת הביטחון של הטווח: {range_info.get('confidence') or 'לא צוינה'}. {range_info.get('confidence_why') or ''}"
    )
    confidence = "medium" if len(fields["missing"]) <= 4 else "low"
    return labels[position], why, action, confidence, position


def _store_offer(conn, document_id: str, text: str, fields: dict, created_at: str) -> str:
    existing = conn.execute("SELECT id FROM purchase_offers WHERE document_id = ?", (document_id,)).fetchone()
    if existing:
        return existing["id"]
    tags = tuple(sorted(item["tag_norm"] for item in fields["assets"]))
    family_key = f"{(fields['buyer'] or '').casefold()}|{','.join(tags)}" if fields["buyer"] and tags else document_id
    family_id = "OFF-" + _digest(family_key)
    prior = conn.execute(
        "SELECT id, version_no FROM purchase_offers WHERE family_id = ? ORDER BY version_no DESC LIMIT 1",
        (family_id,),
    ).fetchone()
    version_no = (prior["version_no"] + 1) if prior else 1
    cancels = bool(re.search(r"replaces the previous|cancels the previous|מחליף את ההצעה|מבטל את ההצעה", text, re.IGNORECASE))
    supersedes = prior["id"] if prior and cancels else ""
    assessment = _assess_offer(conn, fields, created_at)
    offer_id = "OFV-" + _digest(f"{family_id}|{document_id}")
    conn.execute(
        """
        INSERT INTO purchase_offers (
          id, family_id, version_no, document_id, buyer, contacts, offer_date, expiry, price_basis,
          amount, currency, vat, payment_terms, inspection, contingencies, dismantling, loading, transport,
          removal_date, exclusions, missing_terms_json, supersedes_id, assessment_json, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            offer_id,
            family_id,
            version_no,
            document_id,
            fields["buyer"],
            fields["contacts"],
            fields["offer_date"],
            fields["expiry"],
            fields["price_basis"],
            fields["amount"],
            fields["currency"],
            fields["vat"],
            fields["payment_terms"],
            fields["inspection"],
            fields["contingencies"],
            fields["dismantling"],
            fields["loading"],
            fields["transport"],
            fields["removal_date"],
            fields["exclusions"],
            json.dumps(fields["missing"], ensure_ascii=False),
            supersedes,
            json.dumps(assessment, ensure_ascii=False),
            created_at,
        ),
    )
    for item in fields["assets"]:
        conn.execute(
            """
            INSERT OR IGNORE INTO offer_assets (id, offer_id, inventory_row_id, tag_norm, tag_original, role, quantity)
            VALUES (?, ?, ?, ?, ?, 'included', '')
            """,
            ("OA-" + _digest(f"{offer_id}|{item['tag_norm']}|included"), offer_id, item["row_id"], item["tag_norm"], item["tag"],),
        )
    for tag in fields["excluded_tags"]:
        conn.execute(
            """
            INSERT OR IGNORE INTO offer_assets (id, offer_id, inventory_row_id, tag_norm, tag_original, role, quantity)
            VALUES (?, ?, '', ?, ?, 'excluded', '')
            """,
            ("OA-" + _digest(f"{offer_id}|{tag}|excluded"), offer_id, tag, tag),
        )
    if assessment["conclusion"] == "אין די מידע להערכה":
        upsert_review(
            conn,
            kind="offer_gap",
            queue="intake",
            dedupe_key=f"intake:{offer_id}:gap",
            question=assessment["confidence_why"],
            priority="high",
            capture_id=None,
            row_ids=[item["row_id"] for item in fields["assets"] if item.get("row_id")],
            payload={
                "document_id": document_id,
                "offer_id": offer_id,
                "resolve_by": assessment["recommendation"]["label"] + ". " + ", ".join(fields["missing"]),
            },
            created_at=created_at,
        )
    return offer_id


def _store_quote(conn, document_id: str, text: str, created_at: str) -> str:
    existing = conn.execute("SELECT id FROM service_quotes WHERE document_id = ?", (document_id,)).fetchone()
    if existing:
        return existing["id"]
    amounts = _find_money(text)
    amount = amounts[0]["amount"] if len(amounts) == 1 else None
    currency = amounts[0]["currency"] if amount is not None else ""
    kind = "service"
    haystack = text.casefold()
    for label, phrase in (("dismantling", "dismantl"), ("transport", "transport"), ("repair", "repair"), ("inspection", "inspection"), ("פירוק", "פירוק"), ("הובלה", "הובלה")):
        if phrase in haystack:
            kind = label
            break
    quote_id = "SQ-" + _digest(document_id)
    conn.execute(
        """
        INSERT INTO service_quotes (id, document_id, supplier, service_kind, amount, currency, scope_text, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            quote_id,
            document_id,
            _labeled(text, ("supplier", "ספק", "from")),
            kind,
            amount,
            currency,
            text[:500],
            created_at,
        ),
    )
    return quote_id


def _store_appraisal(conn, document_id: str, text: str, matches: dict, created_at: str) -> None:
    from app.valuation import add_evidence, ensure_row_subject

    amounts = _find_money(text)
    if len(amounts) != 1 or len(matches["confirmed"]) != 1:
        return
    row_id = matches["confirmed"][0]["row_id"]
    subject = ensure_row_subject(conn, row_id, created_at)
    title = _labeled(text, ("title", "appraisal", "שומה")) or f"שומה מתוך {document_id}"
    add_evidence(
        conn,
        subject["id"],
        {
            "title": title[:180],
            "source_document": document_id,
            "price_type": "specialist_opinion",
            "price_amount": amounts[0]["amount"],
            "currency": amounts[0]["currency"],
            "included_services": "unknown",
            "tax_treatment": "unknown",
            "role": "context",
            "limitations": "ראיית שווי מהמסמך. היא לא הצעת רכש, ולא נקבעה כטווח בלי ראיות נוספות מאותו סוג.",
        },
        created_at,
    )


def _summary(name: str, extracted: dict, classification: dict, matches: dict, fields: dict | None, conflicts: list[str]) -> dict:
    assets = [item["tag"] for item in matches["confirmed"]]
    unclear = []
    if classification["confidence"] != "clear":
        unclear.append(classification["reason"])
    if matches["candidates"]:
        unclear.append("יש מועמדים שלא שויכו: " + ", ".join(item.get("tag") or item.get("tag_norm") or "" for item in matches["candidates"][:8]))
    if extracted["note"]:
        unclear.append(extracted["note"])
    if conflicts:
        unclear.append("סתירה נשמרה אצל " + ", ".join(sorted(set(conflicts))))
    amounts = []
    if fields:
        if fields["amount"] is not None:
            amounts.append(f"{fields['amount']:g} {fields['currency']}".strip())
        else:
            amounts.extend(item["text"] for item in fields["amounts_seen"][:6])
            if len(fields["amounts_seen"]) > 1:
                unclear.append("יש יותר מסכום אחד, ולא נבחר סכום יחיד.")
        unclear.extend(f"חסר: {item}" for item in fields["missing"])
    what = TYPE_LABELS.get(classification["doc_type"], classification["doc_type"])
    issuer = ""
    stated_date = ""
    if fields:
        issuer = fields["buyer"]
        stated_date = fields["offer_date"]
    else:
        issuer = _labeled(extracted["text"], ("issued by", "from", "מאת", "supplier", "ספק"))
        stated_date = _dates_after(extracted["text"], ("date", "תאריך", "dated"))
    sources = [{"ref": part["ref"], "note": part["text"][:140]} for part in extracted["parts"][:12]]
    next_action = "לשייך פריט או להבהיר את סוג המסמך."
    if classification["doc_type"] == "offer":
        next_action = "לבדוק את ההצעה מול השווי, בלי לאשר אותה."
    elif matches["confirmed"] and classification["doc_type"] == "equipment":
        next_action = "המידע נוסף לפריטים שזוהו. כדאי לקרוא את הסתירות אם יש."
    lines = [
        f"זה {what}. {classification['reason']}",
        f"המסמך אומר: {(extracted['text'] or 'לא נקרא טקסט.')[:400]}",
    ]
    if issuer or stated_date:
        lines.append(f"מי הוציא: {issuer or 'לא צוין'}. תאריך: {stated_date or 'לא צוין'}.")
    if assets:
        lines.append("פריטים שזוהו: " + ", ".join(assets) + ".")
    if amounts:
        lines.append("סכומים שנקראו: " + ", ".join(amounts) + ".")
    if fields and fields["exclusions"]:
        lines.append("החרגות: " + fields["exclusions"][:240])
    if unclear:
        lines.append("לא ברור: " + " ".join(unclear))
    lines.append("הצעד הבא: " + next_action)
    return {
        "what": what,
        "says": (extracted["text"] or "")[:700],
        "issuer": issuer,
        "date": stated_date,
        "assets": assets,
        "amounts": amounts,
        "conditions": [fields["contingencies"]] if fields and fields["contingencies"] else [],
        "deadlines": [item for item in [fields["expiry"] if fields else "", fields["removal_date"] if fields else ""] if item],
        "exclusions": fields["exclusions"] if fields else "",
        "changes": "נוספה גרסה חדשה בלי למחוק את הקודמת." if fields else "נוספו קטעים לפריטים שזוהו, בלי דריסת ניסוח קודם.",
        "unclear": unclear,
        "next_action": next_action,
        "sources": sources,
        "unprocessed": extracted["note"],
        "text": "\n".join(lines),
    }


def _status_for(classification: dict, matches: dict, offer_id: str, assessment: dict | None) -> str:
    if classification["confidence"] != "clear":
        return "awaiting_clarification"
    if classification["doc_type"] == "offer":
        if assessment and assessment.get("position"):
            return "offer_ready"
        return "offer_awaiting_valuation"
    if matches["confirmed"]:
        return "filed"
    return "awaiting_match"


def _public_document(conn, document_id: str) -> dict:
    row = conn.execute("SELECT * FROM intake_documents WHERE id = ?", (document_id,)).fetchone()
    if not row:
        raise LookupError("המסמך לא נמצא.")
    links = [dict(item) for item in conn.execute("SELECT * FROM intake_links WHERE document_id = ?", (document_id,))]
    facts = [dict(item) for item in conn.execute("SELECT * FROM intake_facts WHERE document_id = ?", (document_id,))]
    offer = conn.execute("SELECT * FROM purchase_offers WHERE document_id = ?", (document_id,)).fetchone()
    quote = conn.execute("SELECT * FROM service_quotes WHERE document_id = ?", (document_id,)).fetchone()
    machines = membership_for_tags(conn, [link["tag_norm"] for link in links if link["link_status"] == "confirmed"])
    return {
        "id": row["id"],
        "original_name": row["original_name"],
        "media_kind": row["media_kind"],
        "doc_type": row["doc_type"],
        "type_label": TYPE_LABELS.get(row["doc_type"], row["doc_type"]),
        "type_confidence": row["type_confidence"],
        "classification_reason": row["classification_reason"] or "",
        "status": row["status"],
        "status_label": STATUS_LABELS.get(row["status"], row["status"]),
        "summary": parse_json(row["summary_json"], {}),
        "extraction_note": row["extraction_note"] or "",
        "links": links,
        "facts": facts,
        "machines": machines,
        "offer_id": offer["id"] if offer else "",
        "quote_id": quote["id"] if quote else "",
        "created_at": row["created_at"],
    }


def _public_offer(conn, offer_id: str) -> dict:
    row = conn.execute("SELECT * FROM purchase_offers WHERE id = ?", (offer_id,)).fetchone()
    if not row:
        raise LookupError("ההצעה לא נמצאה.")
    assets = [dict(item) for item in conn.execute("SELECT * FROM offer_assets WHERE offer_id = ?", (offer_id,))]
    versions = [
        {"id": item["id"], "version_no": item["version_no"], "document_id": item["document_id"], "amount": item["amount"], "currency": item["currency"], "created_at": item["created_at"], "supersedes_id": item["supersedes_id"] or ""}
        for item in conn.execute("SELECT * FROM purchase_offers WHERE family_id = ? ORDER BY version_no", (row["family_id"],))
    ]
    assessment = parse_json(row["assessment_json"], {})
    return {
        "id": row["id"],
        "family_id": row["family_id"],
        "version_no": row["version_no"],
        "document_id": row["document_id"],
        "buyer": row["buyer"] or "",
        "contacts": row["contacts"] or "",
        "offer_date": row["offer_date"] or "",
        "expiry": row["expiry"] or "",
        "price_basis": row["price_basis"],
        "amount": row["amount"],
        "currency": row["currency"] or "",
        "vat": row["vat"] or "unknown",
        "payment_terms": row["payment_terms"] or "",
        "inspection": row["inspection"] or "",
        "contingencies": row["contingencies"] or "",
        "dismantling": row["dismantling"] or "",
        "loading": row["loading"] or "",
        "transport": row["transport"] or "",
        "removal_date": row["removal_date"] or "",
        "exclusions": row["exclusions"] or "",
        "missing": parse_json(row["missing_terms_json"], []),
        "supersedes_id": row["supersedes_id"] or "",
        "assets": assets,
        "versions": versions,
        "assessment": assessment,
        "created_at": row["created_at"],
    }


def list_documents(conn, status: str = "") -> list[dict]:
    query = "SELECT id FROM intake_documents"
    args: tuple = ()
    if status:
        query += " WHERE status = ?"
        args = (status,)
    query += " ORDER BY created_at DESC"
    return [_public_document(conn, row["id"]) for row in conn.execute(query, args)]


def list_offers(conn) -> list[dict]:
    rows = conn.execute("SELECT id FROM purchase_offers ORDER BY created_at DESC").fetchall()
    return [_public_offer(conn, row["id"]) for row in rows]


def ingest_file(conn, *, original_name: str, data: bytes, photo_dir: Path, created_at: str) -> dict:
    name = original_name or "document"
    suffix = Path(name).suffix.lower()
    allowed = suffix in IMAGE_SUFFIXES or suffix in DOCUMENT_SUFFIXES or suffix in {".csv", ".xlsx", ".xls", ".pdf", ".doc", ".docx", ".txt", ".rtf"}
    if not data or not allowed:
        return {
            "saved": False,
            "duplicate": False,
            "status": "failed",
            "message": "הקובץ לא נשמר. הוא ריק או מסוג שלא נקרא.",
            "document": None,
        }
    digest = hashlib.sha256(data).hexdigest()
    existing = conn.execute("SELECT id FROM intake_documents WHERE sha256 = ?", (digest,)).fetchone()
    if existing:
        document = _public_document(conn, existing["id"])
        offer = _public_offer(conn, document["offer_id"]) if document["offer_id"] else None
        return {
            "saved": True,
            "duplicate": True,
            "status": document["status"],
            "message": "המסמך כבר נקלט. לא נוצרו פריט, הצעה או משימה נוספים.",
            "document": document,
            "offer": offer,
        }
    document_id = "DOC-" + digest[:16]
    photo_dir.mkdir(parents=True, exist_ok=True)
    stored = photo_dir / f"{document_id}{suffix or '.bin'}"
    stored.write_bytes(data)
    extracted = extract_document(stored)
    classification = classify_text(name, extracted["text"])
    matches = _match_assets(conn, extracted["text"])
    fields = None
    offer_id = ""
    assessment = None
    conflicts: list[str] = []
    if classification["doc_type"] in {"equipment", "mixed", "unclear", "appraisal", "service_quote"}:
        conflicts = _file_equipment(conn, document_id, extracted["parts"], matches, created_at)
    if classification["doc_type"] == "offer":
        fields = _offer_fields(extracted["text"], matches)
        for item in fields["assets"]:
            _link(conn, document_id, item["row_id"], item["tag_norm"], "confirmed", "הפריט נכלל בהצעת הרכש.")
        for item in matches["candidates"]:
            _link(conn, document_id, item.get("row_id"), item.get("tag_norm") or "", "candidate", item.get("reason") or "")
        offer_id = _store_offer(conn, document_id, extracted["text"], fields, created_at)
        assessment = parse_json(conn.execute("SELECT assessment_json FROM purchase_offers WHERE id = ?", (offer_id,)).fetchone()["assessment_json"], {})
    elif classification["doc_type"] == "service_quote":
        _store_quote(conn, document_id, extracted["text"], created_at)
    elif classification["doc_type"] == "appraisal":
        _store_appraisal(conn, document_id, extracted["text"], matches, created_at)
    if classification["confidence"] != "clear":
        upsert_review(
            conn,
            kind="classification",
            queue="intake",
            dedupe_key=f"intake:{document_id}:type",
            question=f"מהו סוג המסמך {name}? {classification['reason']}",
            priority="high",
            capture_id=None,
            row_ids=[item["row_id"] for item in matches["confirmed"]],
            payload={
                "document_id": document_id,
                "resolve_by": "לבחור סוג אחד: מידע על ציוד, הצעת רכש, הצעת שירות, או ראיית שווי.",
            },
            created_at=created_at,
        )
    if matches["candidates"]:
        upsert_review(
            conn,
            kind="asset_match",
            queue="intake",
            dedupe_key=f"intake:{document_id}:match",
            question=f"לאיזו שורה לשייך את {name}? יש יותר ממועמד אחד, ולא נוצר פריט חדש.",
            priority="medium",
            capture_id=None,
            row_ids=[item.get("row_id") for item in matches["candidates"] if item.get("row_id")],
            payload={"document_id": document_id, "candidates": matches["candidates"][:8], "resolve_by": "בחירת שורה אחת או יותר מתוך המועמדים."},
            created_at=created_at,
        )
    summary = _summary(name, extracted, classification, matches, fields, conflicts)
    status = _status_for(classification, matches, offer_id, assessment)
    media = "image" if suffix in IMAGE_SUFFIXES else "spreadsheet" if suffix in {".xlsx", ".xls", ".csv"} else "pdf" if suffix == ".pdf" else "word" if suffix in {".doc", ".docx"} else "text"
    conn.execute(
        """
        INSERT INTO intake_documents (
          id, sha256, original_name, stored_path, media_kind, extracted_text, extraction_note,
          doc_type, type_confidence, classification_reason, summary_json, status, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            document_id,
            digest,
            name,
            str(stored),
            media,
            extracted["text"][:30000],
            extracted["note"],
            classification["doc_type"],
            classification["confidence"],
            classification["reason"],
            json.dumps(summary, ensure_ascii=False),
            status,
            created_at,
            created_at,
        ),
    )
    document = _public_document(conn, document_id)
    offer = _public_offer(conn, offer_id) if offer_id else None
    return {
        "saved": True,
        "duplicate": False,
        "status": status,
        "message": summary["text"],
        "document": document,
        "offer": offer,
    }


def set_type(conn, document_id: str, doc_type: str, created_at: str) -> dict:
    if doc_type not in TYPE_LABELS or doc_type == "mixed":
        raise ValueError("צריך לבחור סוג אחד: מידע על ציוד, הצעת רכש, הצעת שירות, ראיית שווי, או לא ברור.")
    row = conn.execute("SELECT * FROM intake_documents WHERE id = ?", (document_id,)).fetchone()
    if not row:
        raise LookupError("המסמך לא נמצא.")
    extracted = {"text": row["extracted_text"] or "", "parts": [{"ref": "המסמך", "text": row["extracted_text"] or ""}], "note": row["extraction_note"] or ""}
    matches = _match_assets(conn, extracted["text"])
    classification = {"doc_type": doc_type, "confidence": "clear" if doc_type not in {"unclear", "mixed"} else "uncertain", "reason": "הסוג תוקן ידנית."}
    fields = None
    assessment = None
    offer_id = ""
    if doc_type == "offer":
        fields = _offer_fields(extracted["text"], matches)
        offer_id = _store_offer(conn, document_id, extracted["text"], fields, created_at)
        assessment = parse_json(conn.execute("SELECT assessment_json FROM purchase_offers WHERE id = ?", (offer_id,)).fetchone()["assessment_json"], {})
    elif doc_type == "service_quote":
        _store_quote(conn, document_id, extracted["text"], created_at)
    elif doc_type == "equipment":
        _file_equipment(conn, document_id, extracted["parts"], matches, created_at)
    elif doc_type == "appraisal":
        _file_equipment(conn, document_id, extracted["parts"], matches, created_at)
        _store_appraisal(conn, document_id, extracted["text"], matches, created_at)
    status = _status_for(classification, matches, offer_id, assessment)
    summary = parse_json(row["summary_json"], {})
    summary["what"] = TYPE_LABELS[doc_type]
    summary["text"] = f"הסוג תוקן ל{TYPE_LABELS[doc_type]}.\n" + (summary.get("text") or "")
    conn.execute(
        """
        UPDATE intake_documents
        SET doc_type = ?, type_confidence = ?, classification_reason = ?, status = ?, summary_json = ?, updated_at = ?
        WHERE id = ?
        """,
        (doc_type, classification["confidence"], classification["reason"], status, json.dumps(summary, ensure_ascii=False), created_at, document_id),
    )
    if doc_type not in {"unclear", "mixed"}:
        conn.execute(
            "UPDATE review_items SET status = 'resolved', resolution = ?, updated_at = ? WHERE dedupe_key = ? AND status != 'resolved'",
            (f"הסוג תוקן ל{TYPE_LABELS[doc_type]}.", created_at, f"intake:{document_id}:type"),
        )
    return _public_document(conn, document_id)


def set_links(conn, document_id: str, body: dict, created_at: str) -> dict:
    row = conn.execute("SELECT id, extracted_text FROM intake_documents WHERE id = ?", (document_id,)).fetchone()
    if not row:
        raise LookupError("המסמך לא נמצא.")
    action = body.get("action")
    row_ids = [item for item in (body.get("row_ids") or []) if item]
    if action == "remove":
        for row_id in row_ids:
            conn.execute(
                "UPDATE intake_links SET link_status = 'removed', reason = reason || ' הוסר ידנית. שורת המלאי נשארה.' WHERE document_id = ? AND inventory_row_id = ?",
                (document_id, row_id),
            )
    elif action == "add":
        parts = [{"ref": "שיוך ידני", "text": row["extracted_text"] or ""}]
        for row_id in row_ids:
            record = conn.execute("SELECT id, tag_norm, tag_original FROM inventory_rows WHERE id = ? AND kind = 'equipment'", (row_id,)).fetchone()
            if not record:
                continue
            _link(conn, document_id, record["id"], record["tag_norm"] or record["id"], "confirmed", "שויך ידנית. לא נוצרה שורת מלאי.")
            if record["tag_norm"] and record["tag_norm"] in (row["extracted_text"] or "").upper():
                _remember_fact(conn, document_id, record["id"], record["tag_norm"], "excerpt", (row["extracted_text"] or "")[:500], "שיוך ידני", created_at)
        conn.execute(
            "UPDATE review_items SET status = 'resolved', resolution = ?, updated_at = ? WHERE dedupe_key = ? AND status = 'open'",
            ("השיוך עודכן ידנית.", created_at, f"intake:{document_id}:match"),
        )
    else:
        raise ValueError("צריך להוסיף או להסיר שיוך.")
    confirmed = conn.execute(
        "SELECT id FROM intake_links WHERE document_id = ? AND link_status = 'confirmed'",
        (document_id,),
    ).fetchone()
    current = conn.execute("SELECT doc_type, status FROM intake_documents WHERE id = ?", (document_id,)).fetchone()
    if current["doc_type"] == "equipment":
        status = "filed" if confirmed else "awaiting_match"
        conn.execute("UPDATE intake_documents SET status = ?, updated_at = ? WHERE id = ?", (status, created_at, document_id))
    return _public_document(conn, document_id)
