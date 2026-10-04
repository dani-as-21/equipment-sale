"""Capture sync, documentation status, review actions, and export."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import sqlite3
import tempfile
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, UnidentifiedImageError

from app.matching import area_conflict, match_inventory, norm_area, norm_id

ID_RE = re.compile(r"^[0-9a-fA-F-]{36}$")
ITEM_ROLES = {"", "overall", "tag", "nameplate", "accessories", "document"}
GENERAL_ROLES = {"location", "system"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".heic", ".heif", ".tif", ".tiff", ".bmp"}
DOCUMENT_SUFFIXES = {".pdf", ".doc", ".docx", ".txt", ".csv", ".xls", ".xlsx", ".rtf"}
AREA_LABEL_RE = re.compile(
    r"(?<![\w])(?:אזור|area|location)\s*[:：\-]?\s*([^\n\r,;|]{1,60})",
    re.IGNORECASE,
)
ACTOR_USER = "מנהלת הפרויקט"
ACTOR_SYSTEM = "מערכת"


def now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def parse_json(value: str | None, fallback):
    if not value:
        return fallback
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return fallback


def valid_id(value: str) -> bool:
    return bool(value and ID_RE.fullmatch(value))


def link_id(capture_id: str, row_id: str) -> str:
    import hashlib

    digest = hashlib.sha256(f"{capture_id}|{row_id}".encode()).hexdigest()
    return "LNK-" + digest[:16]


def log_change(
    conn: sqlite3.Connection,
    *,
    actor: str,
    action: str,
    entity_type: str,
    entity_id: str,
    prior,
    new,
    created_at: str | None = None,
) -> int:
    cursor = conn.execute(
        """
        INSERT INTO change_log (actor, action, entity_type, entity_id, prior_json, new_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            actor,
            action,
            entity_type,
            entity_id,
            json.dumps(prior, ensure_ascii=False) if prior is not None else None,
            json.dumps(new, ensure_ascii=False) if new is not None else None,
            created_at or now_iso(),
        ),
    )
    return int(cursor.lastrowid)


def equipment_for_match(conn: sqlite3.Connection) -> list[dict]:
    rows = []
    for record in conn.execute("SELECT * FROM inventory_rows WHERE kind = 'equipment'"):
        rows.append(
            {
                "id": record["id"],
                "kind": record["kind"],
                "tag_norm": record["tag_norm"] or "",
                "tag_original": record["tag_original"] or "",
                "serial_norm": record["serial_norm"] or "",
                "serial": record["serial"] or "",
                "model": record["model"] or "",
                "description": record["description"] or "",
                "listed_area": record["listed_area"] or "",
                "identity_area": record["identity_area"] or "",
                "duplicate_group": record["duplicate_group"] or "",
                "sheet_name": record["sheet_name"],
                "original_row": record["original_row"],
            }
        )
    return rows


def area_context(conn: sqlite3.Connection, area_id: str | None, area_name: str | None) -> tuple[str, str, str | None]:
    if area_id:
        area = conn.execute("SELECT * FROM areas WHERE id = ?", (area_id,)).fetchone()
        if area:
            parent = ""
            if area["parent_id"]:
                parent_row = conn.execute("SELECT name FROM areas WHERE id = ?", (area["parent_id"],)).fetchone()
                parent = parent_row["name"] if parent_row else ""
            return area["id"], area["name"], parent
    name = (area_name or "").strip()
    if not name:
        return "", "", ""
    found = conn.execute("SELECT * FROM areas WHERE name = ?", (name,)).fetchone()
    if not found:
        return "", name, ""
    parent = ""
    if found["parent_id"]:
        parent_row = conn.execute("SELECT name FROM areas WHERE id = ?", (found["parent_id"],)).fetchone()
        parent = parent_row["name"] if parent_row else ""
    return found["id"], found["name"], parent


def capture_links(conn: sqlite3.Connection, capture_id: str) -> list[dict]:
    return [
        dict(row)
        for row in conn.execute(
            "SELECT * FROM evidence_links WHERE capture_id = ? ORDER BY created_at",
            (capture_id,),
        )
    ]


def replace_links(
    conn: sqlite3.Connection,
    capture_id: str,
    links: list[dict],
    *,
    actor: str,
    created_at: str,
) -> None:
    prior = capture_links(conn, capture_id)
    prior_key = sorted((item["inventory_row_id"], item["review_status"], item["reason"]) for item in prior)
    new_key = sorted((item["row_id"], item["review_status"], item["reason"]) for item in links)
    if prior_key == new_key:
        return
    conn.execute("DELETE FROM evidence_links WHERE capture_id = ?", (capture_id,))
    for link in links:
        conn.execute(
            """
            INSERT INTO evidence_links (
              id, capture_id, inventory_row_id, method, reason, review_status, creator, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                link_id(capture_id, link["row_id"]),
                capture_id,
                link["row_id"],
                link["method"],
                link["reason"],
                link["review_status"],
                actor,
                created_at,
            ),
        )
    log_change(
        conn,
        actor=actor,
        action="links",
        entity_type="capture",
        entity_id=capture_id,
        prior=prior,
        new=links,
        created_at=created_at,
    )


def linked_row_dicts(conn: sqlite3.Connection, capture_id: str) -> list[dict]:
    return [
        dict(row)
        for row in conn.execute(
            """
            SELECT r.* FROM inventory_rows r
            JOIN evidence_links l ON l.inventory_row_id = r.id
            WHERE l.capture_id = ? AND l.review_status != 'rejected'
            """,
            (capture_id,),
        )
    ]


def maybe_assign_shows(conn: sqlite3.Connection, capture_id: str) -> None:
    confirmed = [
        row
        for row in conn.execute(
            """
            SELECT r.*, l.review_status FROM inventory_rows r
            JOIN evidence_links l ON l.inventory_row_id = r.id
            WHERE l.capture_id = ? AND l.review_status IN ('auto_linked', 'user_confirmed')
            """,
            (capture_id,),
        )
    ]
    if not confirmed:
        return
    groups = {(row["duplicate_group"] or "") for row in confirmed}
    same_item = len(confirmed) == 1 or (len(groups) == 1 and next(iter(groups)))
    if not same_item:
        return
    row_ids = [row["id"] for row in confirmed]
    for photo in conn.execute("SELECT * FROM photos WHERE capture_id = ?", (capture_id,)):
        shows = parse_json(photo["shows_json"], {"scope": "unassigned", "row_ids": []})
        if shows.get("scope") not in (None, "", "unassigned"):
            continue
        role = photo["role"] or ""
        if role in GENERAL_ROLES:
            shows = {"scope": "general", "row_ids": []}
        else:
            shows = {"scope": "rows", "row_ids": row_ids}
        conn.execute("UPDATE photos SET shows_json = ? WHERE id = ?", (json.dumps(shows), photo["id"]))


def upsert_review(
    conn: sqlite3.Connection,
    *,
    kind: str,
    queue: str,
    dedupe_key: str,
    question: str,
    priority: str,
    capture_id: str | None,
    row_ids: list[str],
    payload: dict,
    created_at: str,
) -> None:
    import hashlib

    existing = conn.execute("SELECT id, status FROM review_items WHERE dedupe_key = ?", (dedupe_key,)).fetchone()
    if existing:
        if existing["status"] == "resolved":
            return
        conn.execute(
            """
            UPDATE review_items
            SET question = ?, row_ids_json = ?, payload_json = ?, updated_at = ?, capture_id = COALESCE(?, capture_id)
            WHERE id = ?
            """,
            (
                question,
                json.dumps(row_ids),
                json.dumps(payload, ensure_ascii=False),
                created_at,
                capture_id,
                existing["id"],
            ),
        )
        return
    review_id = "REV-" + hashlib.sha256(dedupe_key.encode()).hexdigest()[:16]
    conn.execute(
        """
        INSERT INTO review_items (
          id, kind, queue, dedupe_key, question, priority, status, capture_id,
          row_ids_json, payload_json, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, 'open', ?, ?, ?, ?, ?)
        """,
        (
            review_id,
            kind,
            queue,
            dedupe_key,
            question,
            priority,
            capture_id,
            json.dumps(row_ids),
            json.dumps(payload, ensure_ascii=False),
            created_at,
            created_at,
        ),
    )


def close_review(conn: sqlite3.Connection, dedupe_key: str, resolution: str, created_at: str) -> None:
    conn.execute(
        """
        UPDATE review_items
        SET status = 'resolved', resolution = ?, updated_at = ?, missing_reason = NULL
        WHERE dedupe_key = ? AND status != 'resolved'
        """,
        (resolution, created_at, dedupe_key),
    )


def apply_capture(conn: sqlite3.Connection, body: dict, created_at: str | None = None) -> dict:
    capture_id = body.get("id") or ""
    if not valid_id(capture_id):
        raise ValueError("מזהה הקליטה אינו תקין.")
    created_at = created_at or now_iso()
    incoming_time = body.get("client_updated_at") or created_at
    existing = conn.execute("SELECT * FROM captures WHERE id = ?", (capture_id,)).fetchone()
    if existing and existing["client_updated_at"] and incoming_time < existing["client_updated_at"]:
        return {"capture": dict(existing), "links": capture_links(conn, capture_id), "kept": "server"}

    area_key, area_name, parent = area_context(conn, body.get("observed_area_id"), body.get("observed_area_name"))
    if body.get("observed_area_name") and not area_name:
        area_name = body.get("observed_area_name")
    visit_id = body.get("visit_id")
    if visit_id and not conn.execute("SELECT id FROM visits WHERE id = ?", (visit_id,)).fetchone():
        visit_id = None
    if not visit_id:
        current = conn.execute("SELECT id FROM visits WHERE is_current = 1").fetchone()
        visit_id = current["id"] if current else None

    prior_fields = dict(existing) if existing else None
    fields = {
        "visit_id": visit_id,
        "observed_area_id": area_key or None,
        "observed_area_name": area_name or body.get("observed_area_name") or "",
        "note": (body.get("note") or "")[:4000],
        "info_provided_by": (body.get("info_provided_by") or "")[:200],
        "tag_text": (body.get("tag_text") or "")[:200],
        "model_text": (body.get("model_text") or "")[:200],
        "serial_text": (body.get("serial_text") or "")[:200],
        "manufacturer_text": (body.get("manufacturer_text") or "")[:200],
        "raw_ocr": body.get("raw_ocr") or "",
        "match_mode": body.get("match_mode") or "none",
        "finalized": 1 if body.get("finalized") else 0,
        "client_updated_at": incoming_time,
        "sync_status": "pending_sync",
    }
    if existing:
        conn.execute(
            """
            UPDATE captures SET
              visit_id = :visit_id,
              observed_area_id = :observed_area_id,
              observed_area_name = :observed_area_name,
              note = :note,
              info_provided_by = :info_provided_by,
              tag_text = :tag_text,
              model_text = :model_text,
              serial_text = :serial_text,
              manufacturer_text = :manufacturer_text,
              raw_ocr = :raw_ocr,
              match_mode = :match_mode,
              finalized = :finalized,
              client_updated_at = :client_updated_at,
              sync_status = :sync_status
            WHERE id = :id
            """,
            {**fields, "id": capture_id},
        )
    else:
        conn.execute(
            """
            INSERT INTO captures (
              id, visit_id, observed_area_id, observed_area_name, note, info_provided_by,
              tag_text, model_text, serial_text, manufacturer_text, raw_ocr, match_mode,
              finalized, client_updated_at, created_at, sync_status
            ) VALUES (
              :id, :visit_id, :observed_area_id, :observed_area_name, :note, :info_provided_by,
              :tag_text, :model_text, :serial_text, :manufacturer_text, :raw_ocr, :match_mode,
              :finalized, :client_updated_at, :created_at, :sync_status
            )
            """,
            {**fields, "id": capture_id, "created_at": body.get("created_at") or created_at},
        )
    if prior_fields and (
        prior_fields.get("tag_text") != fields["tag_text"]
        or prior_fields.get("observed_area_name") != fields["observed_area_name"]
        or prior_fields.get("note") != fields["note"]
    ):
        log_change(
            conn,
            actor=ACTOR_USER,
            action="capture_fields",
            entity_type="capture",
            entity_id=capture_id,
            prior={
                "tag_text": prior_fields.get("tag_text"),
                "observed_area_name": prior_fields.get("observed_area_name"),
                "note": prior_fields.get("note"),
            },
            new={
                "tag_text": fields["tag_text"],
                "observed_area_name": fields["observed_area_name"],
                "note": fields["note"],
            },
            created_at=created_at,
        )

    mode = fields["match_mode"]
    rows = equipment_for_match(conn)
    created_specific_review = False
    if mode == "user":
        user_links = []
        for link in body.get("links") or []:
            row_id = link.get("row_id")
            if not row_id:
                continue
            if not conn.execute("SELECT id FROM inventory_rows WHERE id = ?", (row_id,)).fetchone():
                continue
            user_links.append(
                {
                    "row_id": row_id,
                    "method": "user",
                    "review_status": "user_confirmed",
                    "reason": link.get("reason") or "נבחר ידנית בסיור.",
                }
            )
        replace_links(conn, capture_id, user_links, actor=ACTOR_USER, created_at=created_at)
        maybe_assign_shows(conn, capture_id)
        close_review(conn, f"cap:{capture_id}:unmatched", "קושר ידנית לשורה.", created_at)
    elif mode == "unresolved":
        replace_links(conn, capture_id, [], actor=ACTOR_USER, created_at=created_at)
    else:
        result = match_inventory(
            rows,
            tag=fields["tag_text"],
            serial=fields["serial_text"],
            model=fields["model_text"],
            observed_area=fields["observed_area_name"],
            observed_parent=parent or "",
        )
        server_links = result["links"]
        actor = ACTOR_SYSTEM if result["mode"] == "auto" else ACTOR_SYSTEM
        replace_links(conn, capture_id, server_links, actor=actor, created_at=created_at)
        if result["mode"] == "auto":
            maybe_assign_shows(conn, capture_id)
        fields["match_mode"] = result["mode"] if result["mode"] != "none" else mode
        conn.execute("UPDATE captures SET match_mode = ? WHERE id = ?", (fields["match_mode"], capture_id))
        if body.get("finalized") and result.get("review_kind"):
            created_specific_review = True
            upsert_review(
                conn,
                kind=result["review_kind"],
                queue="visit",
                dedupe_key=f"cap:{capture_id}:{result['review_kind']}",
                question=result["explanation"],
                priority="high",
                capture_id=capture_id,
                row_ids=[item["row_id"] for item in result["suggestions"]],
                payload={"suggestions": result["suggestions"], "explanation": result["explanation"]},
                created_at=created_at,
            )
        elif body.get("finalized") and result["mode"] == "auto":
            close_review(conn, f"cap:{capture_id}:unmatched", "קושר אוטומטית.", created_at)

    if body.get("finalized") and mode == "unresolved":
        upsert_review(
            conn,
            kind="unmatched",
            queue="visit",
            dedupe_key=f"cap:{capture_id}:unmatched",
            question=(
                "הפריט שצולם לא הותאם לשורת מלאי. אפשר לחפש שוב, לקשר ידנית, "
                "או להשאיר אותו כנצפה בסיור. זה לא מכריז על ציוד חדש."
            ),
            priority="high",
            capture_id=capture_id,
            row_ids=[],
            payload={},
            created_at=created_at,
        )
    elif body.get("finalized") and mode not in {"user", "unresolved"} and not created_specific_review:
        current_links = capture_links(conn, capture_id)
        if not current_links:
            upsert_review(
                conn,
                kind="unmatched",
                queue="visit",
                dedupe_key=f"cap:{capture_id}:unmatched",
                question=(
                    "הפריט שצולם לא הותאם לשורת מלאי. אפשר לחפש שוב, לקשר ידנית, "
                    "או להשאיר אותו כנצפה בסיור. זה לא מכריז על ציוד חדש."
                ),
                priority="high",
                capture_id=capture_id,
                row_ids=[],
                payload={},
                created_at=created_at,
            )

    stored = dict(conn.execute("SELECT * FROM captures WHERE id = ?", (capture_id,)).fetchone())
    return {"capture": stored, "links": capture_links(conn, capture_id), "kept": "client"}


def sniff_image_suffix(data: bytes) -> str:
    try:
        with Image.open(io.BytesIO(data)) as image:
            fmt = (image.format or "").lower()
    except (UnidentifiedImageError, OSError):
        return ""
    return {"jpeg": ".jpg", "png": ".png", "webp": ".webp", "gif": ".gif", "tiff": ".tif", "bmp": ".bmp"}.get(fmt, "")


def stored_suffix(original_name: str, data: bytes) -> str:
    suffix = Path(original_name or "").suffix.lower()
    if suffix in IMAGE_SUFFIXES or suffix in DOCUMENT_SUFFIXES:
        return suffix
    sniffed = sniff_image_suffix(data)
    return sniffed or ".jpg"


def tour_type_allowed(original_name: str, data: bytes) -> bool:
    suffix = Path(original_name or "").suffix.lower()
    if suffix in IMAGE_SUFFIXES or suffix in DOCUMENT_SUFFIXES:
        return True
    return bool(sniff_image_suffix(data))


def save_photo_file(
    conn: sqlite3.Connection,
    *,
    photo_id: str,
    capture_id: str,
    role: str,
    original_name: str,
    data: bytes,
    shows: dict | None,
    photo_dir: Path,
    thumb_dir: Path,
    created_at: str,
) -> dict:
    if not valid_id(photo_id) or not valid_id(capture_id):
        raise ValueError("מזהה הצילום או הקליטה אינו תקין.")
    if not data:
        raise ValueError("הקובץ ריק ולא נשמר.")
    photo_dir.mkdir(parents=True, exist_ok=True)
    thumb_dir.mkdir(parents=True, exist_ok=True)
    if not conn.execute("SELECT id FROM captures WHERE id = ?", (capture_id,)).fetchone():
        current = conn.execute("SELECT id FROM visits WHERE is_current = 1").fetchone()
        conn.execute(
            """
            INSERT INTO captures (
              id, visit_id, note, match_mode, finalized, created_at, sync_status, client_updated_at
            ) VALUES (?, ?, '', 'none', 0, ?, 'pending_sync', ?)
            """,
            (capture_id, current["id"] if current else None, created_at, created_at),
        )
    suffix = stored_suffix(original_name, data)
    target = photo_dir / f"{photo_id}{suffix}"
    target.write_bytes(data)
    if not target.exists() or target.stat().st_size != len(data):
        raise IOError("שמירת הקובץ נכשלה.")
    thumb_path = thumb_dir / f"{photo_id}.jpg"
    try:
        with Image.open(io.BytesIO(data)) as image:
            image = image.convert("RGB")
            image.thumbnail((480, 480))
            image.save(thumb_path, "JPEG", quality=70)
    except (UnidentifiedImageError, OSError):
        thumb_path = None
    if shows is None:
        shows = {"scope": "unassigned", "row_ids": []}
    existing = conn.execute("SELECT id FROM photos WHERE id = ?", (photo_id,)).fetchone()
    payload = {
        "id": photo_id,
        "capture_id": capture_id,
        "role": role or "",
        "original_name": original_name or "photo.jpg",
        "stored_path": str(target),
        "thumb_path": str(thumb_path) if thumb_path and thumb_path.exists() else None,
        "upload_state": "stored",
        "shows_json": json.dumps(shows, ensure_ascii=False),
        "created_at": created_at,
    }
    if existing:
        conn.execute(
            """
            UPDATE photos SET capture_id = :capture_id, role = :role, original_name = :original_name,
              stored_path = :stored_path, thumb_path = :thumb_path, upload_state = :upload_state,
              shows_json = :shows_json
            WHERE id = :id
            """,
            payload,
        )
    else:
        conn.execute(
            """
            INSERT INTO photos (
              id, capture_id, role, original_name, stored_path, thumb_path, upload_state, shows_json, created_at
            ) VALUES (
              :id, :capture_id, :role, :original_name, :stored_path, :thumb_path, :upload_state, :shows_json, :created_at
            )
            """,
            payload,
        )
    maybe_assign_shows(conn, capture_id)
    return {"stored": True, "photo_id": photo_id, "synced": False}


def confirm_capture(conn: sqlite3.Connection, capture_id: str, photo_ids: list[str], created_at: str) -> dict:
    if not conn.execute("SELECT id FROM captures WHERE id = ?", (capture_id,)).fetchone():
        return {"synced": False, "missing": photo_ids, "error": "הקליטה לא נמצאה בשרת."}
    missing = []
    for photo_id in photo_ids:
        photo = conn.execute("SELECT * FROM photos WHERE id = ? AND capture_id = ?", (photo_id, capture_id)).fetchone()
        if not photo or not photo["stored_path"] or not Path(photo["stored_path"]).exists():
            missing.append(photo_id)
            continue
        if Path(photo["stored_path"]).stat().st_size <= 0:
            missing.append(photo_id)
    if missing:
        conn.execute(
            "UPDATE captures SET sync_status = 'failed', server_received_at = NULL WHERE id = ?",
            (capture_id,),
        )
        return {"synced": False, "missing": missing}
    conn.execute(
        "UPDATE captures SET sync_status = 'synced', server_received_at = ? WHERE id = ?",
        (created_at, capture_id),
    )
    if photo_ids:
        marks = ",".join("?" for _ in photo_ids)
        conn.execute(
            f"UPDATE photos SET upload_state = 'synced' WHERE capture_id = ? AND id IN ({marks})",
            (capture_id, *photo_ids),
        )
    return {"synced": True, "missing": []}


def _bounded_keys(text: str, keys: dict[str, list]) -> list[str]:
    """Exact identifier hits. A following digit or letter keeps a shorter tag from matching inside a longer one."""
    hay = (text or "").upper()
    ident = set("ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789")
    found: list[str] = []
    for key in keys:
        if len(key) < 4:
            continue
        start = 0
        while True:
            index = hay.find(key, start)
            if index < 0:
                break
            before = hay[index - 1] if index else ""
            after_at = index + len(key)
            after = hay[after_at] if after_at < len(hay) else ""
            if (not before or before not in ident) and (not after or after not in ident):
                found.append(key)
                break
            start = index + 1
    return found


def _read_upload_text(path: Path) -> tuple[str, str]:
    """Return extracted text and a short note when extraction did not run."""
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".csv", ".rtf"}:
            return path.read_text(encoding="utf-8", errors="replace")[:20000], ""
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            return "\n".join(parts)[:20000], ""
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            parts: list[str] = []
            try:
                for sheet in book.worksheets:
                    for row in sheet.iter_rows(max_row=80, max_col=12, values_only=True):
                        for value in row:
                            if value is not None:
                                parts.append(str(value))
                            if len(parts) >= 400:
                                break
            finally:
                book.close()
            return "\n".join(parts)[:20000], ""
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "", "טקסט ה-PDF לא חולץ. ההתאמה משתמשת בשם הקובץ בלבד."
            try:
                reader = PdfReader(str(path))
                parts = [(page.extract_text() or "") for page in reader.pages[:12]]
                return "\n".join(parts)[:20000], ""
            except Exception:
                return "", "טקסט ה-PDF לא חולץ. ההתאמה משתמשת בשם הקובץ בלבד."
    except Exception:
        return "", "הטקסט לא חולץ מהמסמך. ההתאמה משתמשת בשם הקובץ בלבד."
    return "", ""


def _area_stated_by_file(conn: sqlite3.Connection, filename: str, body: str) -> str:
    named = {
        norm_area(row["name"]): row["name"]
        for row in conn.execute("SELECT name FROM areas")
        if norm_area(row["name"]) not in {"", "לא ידוע", "unknown"}
    }
    found: list[str] = []
    for match in AREA_LABEL_RE.finditer(f"{filename}\n{body}"):
        raw = match.group(1).strip(" .,:;|")
        piece = re.split(r"[\s_\-]+", raw)[0] if raw else ""
        options = [raw, piece]
        for option in options:
            hit = named.get(norm_area(option))
            if hit and hit not in found:
                found.append(hit)
                break
    if len(found) == 1:
        return found[0]
    return ""


def _row_public_brief(conn: sqlite3.Connection, row_id: str) -> dict | None:
    record = conn.execute("SELECT * FROM inventory_rows WHERE id = ?", (row_id,)).fetchone()
    if not record:
        return None
    return {
        "id": record["id"],
        "tag": record["tag_original"] or "",
        "description": record["description"] or "",
        "sheet_name": record["sheet_name"] or "",
        "original_row": record["original_row"],
        "listed_area": record["listed_area"] or "",
        "tag_norm": record["tag_norm"] or "",
    }


def _clear_capture_row_ids(conn: sqlite3.Connection, capture_id: str) -> list[str]:
    confirmed = list(
        conn.execute(
            """
            SELECT r.id, r.duplicate_group FROM inventory_rows r
            JOIN evidence_links l ON l.inventory_row_id = r.id
            WHERE l.capture_id = ? AND l.review_status IN ('auto_linked', 'user_confirmed')
            """,
            (capture_id,),
        )
    )
    if not confirmed:
        return []
    groups = {(row["duplicate_group"] or "") for row in confirmed}
    same_item = len(confirmed) == 1 or (len(groups) == 1 and next(iter(groups)))
    if not same_item:
        return []
    return [row["id"] for row in confirmed]


def _stored_file_message(saved: bool, filed: bool, rows: list[dict], shows: dict) -> str:
    labels = shows.get("labels") or []
    if filed and rows:
        message = "נשמר בשרת ותויק אל " + " ; ".join(
            f"{row['tag']} · {row['description']} · גיליון {row['sheet_name']} שורה {row['original_row']}" for row in rows
        )
    elif filed and labels:
        message = "נשמר בשרת ותויק אל " + " ; ".join(
            f"{label.get('tag')} ({label.get('prefix')})" if label.get("prefix") else str(label.get("tag"))
            for label in labels
        )
    elif saved:
        message = "נשמר בשרת. אין התאמה ברורה, והקובץ ממתין בתור הבדיקה."
    else:
        return "לא נשמר בשרת."
    if not filed:
        return message
    areas = []
    for label in labels:
        area = label.get("inventory_area") or ""
        if area and area not in areas:
            areas.append(area)
    if not areas and len(rows) == 1 and rows[0].get("listed_area"):
        areas.append(rows[0]["listed_area"])
    if areas:
        message += " האזור שרשום במלאי: " + ", ".join(areas) + ". זה לא מיקום שאומת מהתמונה."
    for machine in shows.get("machines") or []:
        matched = machine.get("matched_tag") or ""
        others = [tag for tag in (machine.get("other_tags") or []) if tag and tag != matched]
        message += f" הרכיב {matched} נמצא בתוך {machine.get('name')}."
        if others:
            message += " זה לא צילום של " + ", ".join(others) + "."
    return message


def describe_stored_file(conn: sqlite3.Connection, photo_id: str) -> dict:
    photo = conn.execute("SELECT * FROM photos WHERE id = ?", (photo_id,)).fetchone()
    if not photo:
        return {"file_id": photo_id, "saved": False, "filed": False, "message": "הקובץ לא נשמר בשרת."}
    path = Path(photo["stored_path"] or "")
    saved = path.exists() and path.stat().st_size > 0
    shows = parse_json(photo["shows_json"], {})
    capture = conn.execute("SELECT * FROM captures WHERE id = ?", (photo["capture_id"],)).fetchone()
    confirmed_ids = {
        link["inventory_row_id"]
        for link in capture_links(conn, photo["capture_id"])
        if link["review_status"] in {"auto_linked", "user_confirmed"}
    }
    show_ids = shows.get("row_ids") or [] if shows.get("scope") == "rows" else []
    row_ids = [row_id for row_id in show_ids if row_id in confirmed_ids]
    if shows.get("scope") == "tags":
        row_ids = []
    rows = [item for item in (_row_public_brief(conn, row_id) for row_id in row_ids) if item]
    label_names = [label.get("tag") for label in (shows.get("labels") or []) if label.get("tag")]
    filed = saved and (bool(rows) or (shows.get("scope") == "tags" and bool(label_names)))
    review = conn.execute(
        """
        SELECT id FROM review_items
        WHERE capture_id = ? AND status = 'open'
        ORDER BY created_at DESC LIMIT 1
        """,
        (photo["capture_id"],),
    ).fetchone()
    observed = (capture["observed_area_name"] if capture else "") or ""
    area_source = shows.get("area_source") or ("unknown" if not observed else "file")
    message = _stored_file_message(saved, filed, rows, shows)
    return {
        "file_id": photo_id,
        "original_name": photo["original_name"] or "",
        "saved": saved,
        "filed": filed,
        "bytes": path.stat().st_size if saved else 0,
        "capture_id": photo["capture_id"],
        "rows": rows,
        "review_id": None if filed or not review else review["id"],
        "observed_area": observed,
        "area_source": area_source,
        "labels": shows.get("labels") or [],
        "machines": shows.get("machines") or [],
        "inventory_area": shows.get("inventory_area") or (rows[0]["listed_area"] if len(rows) == 1 else ""),
        "message": message,
    }


def _remember_shows(conn: sqlite3.Connection, photo_id: str, shows: dict) -> None:
    conn.execute("UPDATE photos SET shows_json = ? WHERE id = ?", (json.dumps(shows, ensure_ascii=False), photo_id))


def _mark_capture_files_synced(conn: sqlite3.Connection, capture_id: str, created_at: str) -> bool:
    photos = list(conn.execute("SELECT * FROM photos WHERE capture_id = ?", (capture_id,)))
    if not photos:
        return False
    for photo in photos:
        path = Path(photo["stored_path"] or "")
        if not path.exists() or path.stat().st_size <= 0:
            conn.execute("UPDATE captures SET sync_status = 'failed', server_received_at = NULL WHERE id = ?", (capture_id,))
            return False
    conn.execute(
        "UPDATE captures SET sync_status = 'synced', server_received_at = ? WHERE id = ?",
        (created_at, capture_id),
    )
    conn.execute("UPDATE photos SET upload_state = 'synced' WHERE capture_id = ?", (capture_id,))
    return True


def file_tour_upload(
    conn: sqlite3.Connection,
    *,
    file_id: str,
    original_name: str,
    data: bytes,
    capture_id: str,
    allow_open_capture: bool,
    photo_dir: Path,
    thumb_dir: Path,
    created_at: str,
) -> dict:
    """Store one tour file on the server and file it only when the match is clear.

    The caller's selected area is not an input. Observed area stays unknown unless
    the file text names one, or the file is attached to an existing capture that
    already has one.
    """
    name = original_name or "file"
    if not valid_id(file_id):
        return {"file_id": file_id, "original_name": name, "saved": False, "filed": False, "rows": [], "message": "הקובץ לא נשמר בשרת."}
    if not data:
        return {"file_id": file_id, "original_name": name, "saved": False, "filed": False, "rows": [], "message": "הקובץ ריק ולא נשמר בשרת."}
    if not tour_type_allowed(name, data):
        return {
            "file_id": file_id,
            "original_name": name,
            "saved": False,
            "filed": False,
            "rows": [],
            "message": "סוג הקובץ לא נתמך, והוא לא נשמר בשרת.",
        }
    existing = conn.execute("SELECT stored_path FROM photos WHERE id = ?", (file_id,)).fetchone()
    if existing and existing["stored_path"] and Path(existing["stored_path"]).exists() and Path(existing["stored_path"]).stat().st_size > 0:
        return describe_stored_file(conn, file_id)

    suffix = stored_suffix(name, data)
    photo_dir.mkdir(parents=True, exist_ok=True)
    handle, tmp_name = tempfile.mkstemp(suffix=suffix, dir=str(photo_dir))
    os.close(handle)
    tmp_path = Path(tmp_name)
    os_note = ""
    body = ""
    ocr_raw = ""
    try:
        tmp_path.write_bytes(data)
        if suffix in IMAGE_SUFFIXES or sniff_image_suffix(data):
            ocr = run_ocr(str(tmp_path))
            ocr_raw = ocr.get("raw_text") or ""
            body = ocr_raw
            if not ocr.get("available"):
                os_note = ocr.get("message") or ""
        else:
            body, os_note = _read_upload_text(tmp_path)
    finally:
        tmp_path.unlink(missing_ok=True)

    equipment = equipment_for_match(conn)
    tags: dict[str, list[dict]] = {}
    serials: dict[str, list[dict]] = {}
    for row in equipment:
        if row.get("tag_norm"):
            tags.setdefault(row["tag_norm"], []).append(row)
        if row.get("serial_norm"):
            serials.setdefault(row["serial_norm"], []).append(row)
    haystack = f"{name}\n{body}"
    exact_tags = _bounded_keys(haystack, tags)
    exact_serials = [token for token in _bounded_keys(haystack, serials) if token not in tags]
    stated_area = _area_stated_by_file(conn, Path(name).stem, body)
    area_id, area_name, parent = ("", "", "")
    area_source = "unknown"
    if stated_area:
        area_id, area_name, parent = area_context(conn, None, stated_area)
        area_source = "file"
        if not area_name:
            area_name = stated_area

    open_rows: list[str] = []
    open_capture = None
    if allow_open_capture and capture_id and valid_id(capture_id):
        open_capture = conn.execute("SELECT * FROM captures WHERE id = ?", (capture_id,)).fetchone()
        if open_capture:
            open_rows = _clear_capture_row_ids(conn, capture_id)

    multi_clear = []
    if len(exact_tags) > 1 and not stated_area:
        multi_clear = []
        blocked = False
        for token in exact_tags:
            probed = match_inventory(equipment, tag=token, observed_area=area_name, observed_parent=parent or "")
            if probed.get("mode") != "auto":
                blocked = True
                break
            multi_clear.extend(probed["links"])
        if blocked:
            multi_clear = []
    identifiers_conflict = False
    if len(exact_tags) == 1 and exact_serials:
        tag_ids = {row["id"] for row in tags[exact_tags[0]]}
        serial_ids = {row["id"] for token in exact_serials for row in serials[token]}
        if serial_ids and tag_ids.isdisjoint(serial_ids):
            identifiers_conflict = True
    chosen_tag = exact_tags[0] if len(exact_tags) == 1 and not identifiers_conflict else ""
    if multi_clear:
        chosen_tag = ""
        identifiers_conflict = False
    chosen_serial = exact_serials[0] if len(exact_serials) == 1 and not exact_tags else ""
    match = None
    if chosen_tag or chosen_serial:
        match = match_inventory(
            equipment,
            tag=chosen_tag,
            serial=chosen_serial,
            observed_area=area_name,
            observed_parent=parent or "",
        )
    suggestions: list[dict] = []
    if len(exact_tags) > 1 or len(exact_serials) > 1 or identifiers_conflict:
        for token in exact_tags:
            for row in tags.get(token, [])[:8]:
                suggestions.append(
                    {
                        "row_id": row["id"],
                        "tag": row.get("tag_original") or "",
                        "description": row.get("description") or "",
                        "sheet_name": row.get("sheet_name") or "",
                        "original_row": row.get("original_row"),
                        "why": "התג מופיע בקובץ יחד עם תג אחר. לא נבחרה שורה.",
                    }
                )
        for token in exact_serials:
            if token in exact_tags:
                continue
            for row in serials.get(token, [])[:8]:
                suggestions.append(
                    {
                        "row_id": row["id"],
                        "tag": row.get("tag_original") or "",
                        "description": row.get("description") or "",
                        "sheet_name": row.get("sheet_name") or "",
                        "original_row": row.get("original_row"),
                        "why": "המספר הסידורי מופיע בקובץ יחד עם מזהה אחר. לא נבחרה שורה.",
                    }
                )
    elif match and match.get("mode") != "auto":
        suggestions = match.get("suggestions") or []
    elif not chosen_tag and not chosen_serial:
        stem = norm_id(Path(name).stem)
        probes = [stem] if stem else []
        for line in (body or "").splitlines():
            line = line.strip()
            if 4 <= len(line) <= 40:
                probes.append(line)
            if len(probes) >= 6:
                break
        for probe in probes:
            probed = match_inventory(equipment, tag=probe, observed_area=area_name, observed_parent=parent or "")
            if probed.get("suggestions") and probed.get("mode") != "auto":
                suggestions = probed["suggestions"]
                match = probed
                break

    use_open = False
    target_capture = ""
    if match and match.get("mode") == "auto":
        target_capture = str(uuid.uuid4())
        if open_rows and {link["row_id"] for link in match["links"]} == set(open_rows):
            target_capture = capture_id
            use_open = True
            if open_capture and (open_capture["observed_area_name"] or "") and area_source != "file":
                area_source = "existing_capture"
                area_name = open_capture["observed_area_name"] or ""
    elif not exact_tags and not exact_serials and not suggestions and open_rows:
        file_area_conflicts = False
        if area_source == "file" and open_capture and (open_capture["observed_area_name"] or ""):
            file_area_conflicts = area_conflict(area_name, open_capture["observed_area_name"] or "", "")
        if not file_area_conflicts:
            use_open = True
            target_capture = capture_id
            area_source = "existing_capture" if (open_capture and (open_capture["observed_area_name"] or "")) else area_source
            if area_source == "existing_capture":
                area_name = open_capture["observed_area_name"] or ""
                area_id = open_capture["observed_area_id"] or ""
    if not target_capture:
        target_capture = str(uuid.uuid4())

    if not use_open:
        filing_tag = chosen_tag if match and match.get("mode") == "auto" else ""
        filing_serial = chosen_serial if match and match.get("mode") == "auto" else ""
        if match and match.get("mode") == "pending" and len(exact_tags) <= 1 and len(exact_serials) <= 1:
            filing_tag = chosen_tag
            filing_serial = chosen_serial
        mode = "auto"
        user_links = None
        if multi_clear:
            mode = "user"
            user_links = []
            seen = set()
            for link in multi_clear:
                if link["row_id"] in seen:
                    continue
                seen.add(link["row_id"])
                user_links.append({**link, "method": "user", "review_status": "user_confirmed", "reason": "כל תג שנראה בקובץ חד-משמעי. רכיב שלא נראה לא סומן."})
            filing_tag = " ".join(exact_tags)
        elif len(exact_tags) > 1 or len(exact_serials) > 1 or identifiers_conflict:
            mode = "unresolved"
            filing_tag = ""
            filing_serial = ""
        payload = {
            "id": target_capture,
            "observed_area_id": area_id or None,
            "observed_area_name": area_name if area_source == "file" else "",
            "tag_text": filing_tag,
            "serial_text": filing_serial,
            "raw_ocr": (ocr_raw or body or "")[:8000],
            "match_mode": mode,
            "finalized": 1,
            "note": "",
            "client_updated_at": created_at,
            "created_at": created_at,
        }
        if user_links is not None:
            payload["links"] = user_links
        apply_capture(conn, payload, created_at)
        if area_source != "file":
            conn.execute(
                "UPDATE captures SET observed_area_id = NULL, observed_area_name = '' WHERE id = ?",
                (target_capture,),
            )

    role = "document" if suffix in DOCUMENT_SUFFIXES and not sniff_image_suffix(data) else "nameplate"
    if not (chosen_tag or chosen_serial or ocr_raw):
        role = "document" if suffix in DOCUMENT_SUFFIXES else "overall"
    try:
        save_photo_file(
            conn,
            photo_id=file_id,
            capture_id=target_capture,
            role=role,
            original_name=name,
            data=data,
            shows={"scope": "unassigned", "row_ids": [], "area_source": area_source},
            photo_dir=photo_dir,
            thumb_dir=thumb_dir,
            created_at=created_at,
        )
    except (ValueError, OSError):
        if not use_open:
            conn.execute("DELETE FROM evidence_links WHERE capture_id = ?", (target_capture,))
            conn.execute("DELETE FROM review_items WHERE capture_id = ?", (target_capture,))
            conn.execute("DELETE FROM photos WHERE capture_id = ?", (target_capture,))
            conn.execute("DELETE FROM captures WHERE id = ?", (target_capture,))
        return {"file_id": file_id, "original_name": name, "saved": False, "filed": False, "rows": [], "message": "הקובץ לא נשמר בשרת."}

    if ocr_raw:
        conn.execute("UPDATE photos SET ocr_raw = ? WHERE id = ?", (ocr_raw[:8000], file_id))
    maybe_assign_shows(conn, target_capture)
    photo = conn.execute("SELECT shows_json FROM photos WHERE id = ?", (file_id,)).fetchone()
    shows = parse_json(photo["shows_json"], {"scope": "unassigned", "row_ids": []})
    if area_source == "file":
        shows["area_source"] = "file"
    elif area_source == "existing_capture":
        shows["area_source"] = "existing_capture"
    else:
        shows["area_source"] = "unknown"
    labels = []
    linked_ids = [link["inventory_row_id"] for link in capture_links(conn, target_capture) if link["review_status"] in {"auto_linked", "user_confirmed"}]
    if linked_ids:
        shows["scope"] = "rows"
        shows["row_ids"] = linked_ids
        for row_id in linked_ids:
            brief = _row_public_brief(conn, row_id)
            if not brief:
                continue
            labels.append({
                "row_id": row_id,
                "tag": brief["tag"],
                "tag_norm": brief.get("tag_norm") or (brief["tag"] or "").upper(),
                "prefix": __import__("app.machines", fromlist=["prefix_label"]).prefix_label(brief["tag"]),
                "inventory_area": brief["listed_area"],
            })
    else:
        from app.machines import membership_for_tags

        member_index = {
            row["tag_norm"]: row["tag_original"]
            for row in conn.execute("SELECT tag_norm, tag_original FROM machine_members WHERE link_status = 'active'")
        }
        seen_members = _bounded_keys(haystack, {tag: [] for tag in member_index if tag not in tags})
        if seen_members and not exact_tags and not exact_serials:
            shows["scope"] = "tags"
            shows["row_ids"] = []
            for token in seen_members:
                labels.append({
                    "row_id": "",
                    "tag": member_index[token],
                    "tag_norm": token,
                    "prefix": __import__("app.machines", fromlist=["prefix_label"]).prefix_label(token),
                    "inventory_area": "",
                })
    shows["labels"] = labels
    if labels:
        from app.machines import membership_for_tags

        shows["machines"] = membership_for_tags(conn, [label["tag_norm"] for label in labels])
        areas = [label["inventory_area"] for label in labels if label.get("inventory_area")]
        shows["inventory_area"] = areas[0] if len(set(areas)) == 1 else ""
    _remember_shows(conn, file_id, shows)
    synced = _mark_capture_files_synced(conn, target_capture, created_at)
    filed_now = (shows.get("scope") == "rows" and bool(shows.get("row_ids"))) or (shows.get("scope") == "tags" and bool(shows.get("labels")))
    if not filed_now:
        review = conn.execute(
            "SELECT id, question FROM review_items WHERE capture_id = ? AND status = 'open' ORDER BY created_at DESC LIMIT 1",
            (target_capture,),
        ).fetchone()
        question = f"הקובץ {name} נשמר בשרת. אין התאמה ברורה לשורת מלאי, ולכן הוא לא שויך בניחוש."
        if match and match.get("explanation"):
            question = f"{question} {match['explanation']}"
        if os_note:
            question = f"{question} {os_note}"
        row_ids = [item.get("row_id") for item in suggestions if item.get("row_id")][:12]
        if review:
            conn.execute(
                """
                UPDATE review_items
                SET question = ?, payload_json = ?, row_ids_json = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    question,
                    json.dumps({
                        "suggestions": suggestions[:12],
                        "filename": name,
                        "photo_ids": [file_id],
                        "resolve_by": "תג שנקרא בבירור, או בחירה מבין השורות המוצגות. אפשר להעלות עוד קבצים לפני שהשאלה נסגרת.",
                    }, ensure_ascii=False),
                    json.dumps(row_ids),
                    created_at,
                    review["id"],
                ),
            )
        else:
            upsert_review(
                conn,
                kind="unmatched",
                queue="visit",
                dedupe_key=f"cap:{target_capture}:unmatched",
                question=question,
                priority="high",
                capture_id=target_capture,
                row_ids=[item.get("row_id") for item in suggestions if item.get("row_id")][:12],
                payload={
                    "suggestions": suggestions[:12],
                    "filename": name,
                    "photo_ids": [file_id],
                    "resolve_by": "תג שנקרא בבירור, או בחירה מבין השורות המוצגות. אפשר להעלות עוד קבצים לפני שהשאלה נסגרת.",
                },
                created_at=created_at,
            )
    result = describe_stored_file(conn, file_id)
    if not result.get("saved") or not synced:
        result["saved"] = bool(result.get("saved"))
    if not result.get("saved"):
        result["filed"] = False
        result["message"] = "לא נשמר בשרת."
    return result


def run_ocr(path: str) -> dict:
    if not path or not Path(path).exists():
        return {"available": False, "raw_text": "", "message": "הקובץ לא נמצא."}
    if not shutil.which("tesseract"):
        return {
            "available": False,
            "raw_text": "",
            "message": "זיהוי הטקסט לא זמין כרגע. אפשר להקליד את התג. הצילום נשמר גם בלי זה.",
        }
    try:
        import pytesseract

        with Image.open(path) as image:
            image.thumbnail((1800, 1800))
            text = pytesseract.image_to_string(image, lang="eng") or ""
        return {"available": True, "raw_text": text, "message": ""}
    except Exception:
        return {
            "available": True,
            "raw_text": "",
            "message": "לא הצלחנו לקרוא טקסט מהתמונה. אפשר להקליד את התג או לסמן שהשלט לא קריא.",
        }


def update_photo(conn: sqlite3.Connection, photo_id: str, body: dict, created_at: str) -> dict:
    photo = conn.execute("SELECT * FROM photos WHERE id = ?", (photo_id,)).fetchone()
    if not photo:
        raise LookupError("הצילום לא נמצא.")
    prior = {"role": photo["role"], "shows_json": photo["shows_json"]}
    role = body.get("role", photo["role"]) or ""
    shows = body.get("shows", parse_json(photo["shows_json"], {"scope": "unassigned", "row_ids": []}))
    conn.execute(
        "UPDATE photos SET role = ?, shows_json = ? WHERE id = ?",
        (role, json.dumps(shows, ensure_ascii=False), photo_id),
    )
    log_change(
        conn,
        actor=ACTOR_USER,
        action="photo_assignment",
        entity_type="photo",
        entity_id=photo_id,
        prior=prior,
        new={"role": role, "shows": shows},
        created_at=created_at,
    )
    return {"id": photo_id, "role": role, "shows": shows}


def split_capture(conn: sqlite3.Connection, capture_id: str, new_id: str, photo_ids: list[str], created_at: str) -> dict:
    if not valid_id(new_id):
        raise ValueError("מזהה הקליטה החדשה אינו תקין.")
    source = conn.execute("SELECT * FROM captures WHERE id = ?", (capture_id,)).fetchone()
    if not source:
        raise LookupError("הקליטה לא נמצאה.")
    if not conn.execute("SELECT id FROM captures WHERE id = ?", (new_id,)).fetchone():
        conn.execute(
            """
            INSERT INTO captures (
              id, visit_id, observed_area_id, observed_area_name, note, match_mode, finalized,
              created_at, sync_status, client_updated_at
            ) VALUES (?, ?, ?, ?, '', 'unresolved', 1, ?, 'pending_sync', ?)
            """,
            (
                new_id,
                source["visit_id"],
                source["observed_area_id"],
                source["observed_area_name"],
                created_at,
                created_at,
            ),
        )
    moved = []
    for photo_id in photo_ids:
        photo = conn.execute(
            "SELECT id FROM photos WHERE id = ? AND capture_id = ?",
            (photo_id, capture_id),
        ).fetchone()
        if not photo:
            continue
        conn.execute("UPDATE photos SET capture_id = ? WHERE id = ?", (new_id, photo_id))
        moved.append(photo_id)
    log_change(
        conn,
        actor=ACTOR_USER,
        action="split",
        entity_type="capture",
        entity_id=capture_id,
        prior={"photo_ids": photo_ids},
        new={"new_capture_id": new_id, "moved": moved},
        created_at=created_at,
    )
    return {"new_capture_id": new_id, "moved": moved}


def add_relationship(conn: sqlite3.Connection, capture_id: str, body: dict, created_at: str) -> dict:
    text = (body.get("text") or "").strip()
    if not text:
        raise ValueError("חסר ניסוח של הקשר.")
    status = body.get("status") or "tentative"
    if status not in {"tentative", "confirmed", "unclear"}:
        status = "tentative"
    claim_id = "REL-" + uuid.uuid4().hex[:16]
    row_ids = body.get("row_ids") or []
    conn.execute(
        """
        INSERT INTO relationship_claims (id, capture_id, text, person, status, related_row_ids_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            claim_id,
            capture_id,
            text[:2000],
            (body.get("person") or "")[:200],
            status,
            json.dumps(row_ids),
            created_at,
        ),
    )
    if status == "unclear":
        upsert_review(
            conn,
            kind="relationship",
            queue="visit",
            dedupe_key=f"cap:{capture_id}:relationship:{claim_id}",
            question=f"נרשמה הערה על קשר, והיא לא ברורה: {text} האם זה נכון, ולגבי איזה ציוד?",
            priority="medium",
            capture_id=capture_id,
            row_ids=row_ids,
            payload={"claim_id": claim_id},
            created_at=created_at,
        )
    return {"id": claim_id, "status": status}


def undo_change(conn: sqlite3.Connection, log_id: int, created_at: str) -> dict:
    entry = conn.execute("SELECT * FROM change_log WHERE id = ?", (log_id,)).fetchone()
    if not entry or entry["undone"]:
        raise LookupError("אין פעולה לביטול.")
    prior = parse_json(entry["prior_json"], None)
    if entry["action"] == "links" and entry["entity_type"] == "capture":
        restored = []
        for link in prior or []:
            restored.append(
                {
                    "row_id": link["inventory_row_id"],
                    "method": link["method"],
                    "review_status": link["review_status"],
                    "reason": link["reason"],
                }
            )
        replace_links(conn, entry["entity_id"], restored, actor=ACTOR_USER, created_at=created_at)
    elif entry["action"] == "capture_fields":
        prior = prior or {}
        conn.execute(
            """
            UPDATE captures SET tag_text = ?, observed_area_name = ?, note = ?, client_updated_at = ?
            WHERE id = ?
            """,
            (prior.get("tag_text"), prior.get("observed_area_name"), prior.get("note"), created_at, entry["entity_id"]),
        )
    elif entry["action"] == "photo_assignment":
        prior = prior or {}
        conn.execute(
            "UPDATE photos SET role = ?, shows_json = ? WHERE id = ?",
            (prior.get("role"), prior.get("shows_json"), entry["entity_id"]),
        )
    else:
        raise ValueError("אי אפשר לבטל את הפעולה הזו אוטומטית.")
    conn.execute("UPDATE change_log SET undone = 1 WHERE id = ?", (log_id,))
    log_change(
        conn,
        actor=ACTOR_USER,
        action="undo",
        entity_type=entry["entity_type"],
        entity_id=entry["entity_id"],
        prior=parse_json(entry["new_json"], None),
        new=prior,
        created_at=created_at,
    )
    return {"undone": log_id}


def attach_review_file(
    conn: sqlite3.Connection,
    review_id: str,
    *,
    file_id: str,
    original_name: str,
    data: bytes,
    photo_dir: Path,
    thumb_dir: Path,
    created_at: str,
) -> dict:
    """Store supporting material on a task and file it when the tag itself is clear."""
    item = conn.execute("SELECT * FROM review_items WHERE id = ?", (review_id,)).fetchone()
    if not item:
        raise LookupError("השאלה לא נמצאה.")
    result = file_tour_upload(
        conn,
        file_id=file_id,
        original_name=original_name,
        data=data,
        capture_id="",
        allow_open_capture=False,
        photo_dir=photo_dir,
        thumb_dir=thumb_dir,
        created_at=created_at,
    )
    payload = parse_json(item["payload_json"], {})
    photo_ids = [str(photo_id) for photo_id in (payload.get("photo_ids") or [])]
    if result.get("saved") and result.get("file_id") and result["file_id"] not in photo_ids:
        photo_ids.append(result["file_id"])
    payload["photo_ids"] = photo_ids
    payload["last_upload"] = {
        "file_id": result.get("file_id"),
        "saved": bool(result.get("saved")),
        "filed": bool(result.get("filed")),
        "message": result.get("message") or "",
        "rows": result.get("rows") or [],
        "labels": result.get("labels") or [],
    }
    if result.get("saved") and result.get("filed"):
        message = (result.get("message") or "נשמר בשרת.") + " הרשומה שהתאמה אליה עודכנה. השאלה נשארת פתוחה עד תשובה."
        missing = None
    elif result.get("saved"):
        message = (result.get("message") or "נשמר בשרת.") + " השאלה נשארת פתוחה."
        missing = message
    else:
        message = result.get("message") or "הקובץ לא נשמר בשרת."
        missing = "הקובץ לא נשמר, והשאלה לא נסגרה."
    conn.execute(
        """
        UPDATE review_items
        SET payload_json = ?, missing_reason = ?, updated_at = ?, status = CASE WHEN status = 'resolved' THEN status ELSE 'open' END
        WHERE id = ?
        """,
        (json.dumps(payload, ensure_ascii=False), missing, created_at, review_id),
    )
    return {
        "ok": bool(result.get("saved")),
        "closed": False,
        "message": message,
        "photo_id": result.get("file_id"),
        "filed": bool(result.get("filed")),
    }


def review_action(conn: sqlite3.Connection, review_id: str, body: dict, created_at: str) -> dict:
    item = conn.execute("SELECT * FROM review_items WHERE id = ?", (review_id,)).fetchone()
    if not item:
        raise LookupError("השאלה לא נמצאה.")
    action = body.get("action")
    if action == "defer":
        conn.execute(
            "UPDATE review_items SET status = 'deferred', updated_at = ?, missing_reason = NULL WHERE id = ?",
            (created_at, review_id),
        )
        return {"status": "deferred"}
    if action == "link":
        row_ids = [row_id for row_id in (body.get("row_ids") or []) if row_id]
        if not item["capture_id"] or not row_ids:
            conn.execute(
                "UPDATE review_items SET missing_reason = ?, updated_at = ? WHERE id = ?",
                ("כדי לסגור צריך לבחור שורה מהרשימה. העלאה או טקסט כללי לא מספיקים.", created_at, review_id),
            )
            return {"status": item["status"], "missing": "כדי לסגור צריך לבחור שורה מהרשימה."}
        links = [
            {"row_id": row_id, "method": "user", "review_status": "user_confirmed", "reason": "נבחר מתוך תור הבדיקה."}
            for row_id in row_ids
        ]
        replace_links(conn, item["capture_id"], links, actor=ACTOR_USER, created_at=created_at)
        conn.execute(
            "UPDATE captures SET match_mode = 'user', client_updated_at = ? WHERE id = ?",
            (created_at, item["capture_id"]),
        )
        maybe_assign_shows(conn, item["capture_id"])
        conn.execute(
            """
            UPDATE review_items
            SET status = 'resolved', resolution = ?, updated_at = ?, missing_reason = NULL
            WHERE id = ?
            """,
            ("קושר לשורות שנבחרו.", created_at, review_id),
        )
        return {"status": "resolved"}
    if action == "leave_separate":
        if item["kind"] == "grouping":
            payload = parse_json(item["payload_json"], {})
            machine = payload.get("machine_id")
            if not machine:
                return {"status": item["status"], "missing": "אין יחידה לפרק."}
            from app.machines import machine_action

            machine_action(conn, machine, {"action": "dissolve"}, created_at)
            return {"status": "resolved"}
        if item["kind"] not in {"duplicate_tag", "generic_name", "ambiguous", "unmatched"}:
            return {"status": item["status"], "missing": "הפעולה הזו לא מתאימה לשאלה."}
        conn.execute(
            """
            UPDATE review_items
            SET status = 'resolved', resolution = ?, updated_at = ?, missing_reason = NULL
            WHERE id = ?
            """,
            ("השורות נשארות נפרדות. לא אוחדו ולא הוכרז שהן מערכת אחת.", created_at, review_id),
        )
        return {"status": "resolved"}
    if action == "correct":
        if not item["capture_id"]:
            return {"status": item["status"], "missing": "אין קליטה לעדכון."}
        capture = conn.execute("SELECT * FROM captures WHERE id = ?", (item["capture_id"],)).fetchone()
        prior = {
            "tag_text": capture["tag_text"],
            "observed_area_name": capture["observed_area_name"],
            "note": capture["note"],
        }
        tag_text = body.get("tag_text", capture["tag_text"])
        area_name = body.get("observed_area_name", capture["observed_area_name"])
        conn.execute(
            """
            UPDATE captures SET tag_text = ?, observed_area_name = ?, client_updated_at = ? WHERE id = ?
            """,
            (tag_text, area_name, created_at, item["capture_id"]),
        )
        log_change(
            conn,
            actor=ACTOR_USER,
            action="capture_fields",
            entity_type="capture",
            entity_id=item["capture_id"],
            prior=prior,
            new={"tag_text": tag_text, "observed_area_name": area_name, "note": capture["note"]},
            created_at=created_at,
        )
        if body.get("recompute"):
            apply_capture(
                conn,
                {
                    "id": item["capture_id"],
                    "client_updated_at": created_at,
                    "tag_text": tag_text,
                    "observed_area_name": area_name,
                    "note": capture["note"],
                    "serial_text": capture["serial_text"],
                    "model_text": capture["model_text"],
                    "match_mode": "auto",
                    "finalized": 1,
                    "info_provided_by": capture["info_provided_by"],
                    "raw_ocr": capture["raw_ocr"],
                    "visit_id": capture["visit_id"],
                },
                created_at,
            )
        return {"status": "open", "corrected": True}
    if action == "resolve" and item["kind"] == "grouping":
        text = (body.get("text") or "").strip()
        if len(text) < 2:
            conn.execute(
                "UPDATE review_items SET missing_reason = ?, updated_at = ? WHERE id = ?",
                ("חסרה תשובה כתובה. העלאת קובץ לבד לא סוגרת את השאלה.", created_at, review_id),
            )
            return {"status": item["status"], "missing": "חסרה תשובה כתובה. העלאת קובץ לבד לא סוגרת את השאלה."}
        payload = parse_json(item["payload_json"], {})
        machine = payload.get("machine_id")
        if machine:
            from app.machines import machine_action

            machine_action(conn, machine, {"action": "note", "text": text}, created_at)
        return {
            "status": "open",
            "noted": True,
            "message": "התשובה נשמרה על היחידה. המכירה יחד נשארת פתוחה עד אישור או פירוק.",
        }
    if action == "resolve" and str(item["kind"]).startswith("valuation"):
        return {
            "status": item["status"],
            "missing": "עונים על שאלת השווי מכרטיס השווי, כדי שהתשובה תתעדכן במוכנות ולא תיכנס כמחיר.",
        }
    if action == "resolve":
        text = (body.get("text") or "").strip()
        if len(text) < 2:
            conn.execute(
                "UPDATE review_items SET missing_reason = ?, updated_at = ? WHERE id = ?",
                ("חסרה תשובה כתובה. העלאת קובץ לבד לא סוגרת את השאלה.", created_at, review_id),
            )
            return {"status": item["status"], "missing": "חסרה תשובה כתובה. העלאת קובץ לבד לא סוגרת את השאלה."}
        if item["kind"] in {"lab_anomaly", "nameplate", "relationship", "generic_name", "unmatched"}:
            conn.execute(
                """
                UPDATE review_items
                SET status = 'resolved', resolution = ?, updated_at = ?, missing_reason = NULL
                WHERE id = ?
                """,
                (text, created_at, review_id),
            )
            return {"status": "resolved"}
        conn.execute(
            "UPDATE review_items SET missing_reason = ?, updated_at = ? WHERE id = ?",
            ("כדי לסגור את השאלה צריך לקשר שורה, להשאיר את השורות נפרדות, או לדחות.", created_at, review_id),
        )
        return {"status": item["status"], "missing": "כדי לסגור את השאלה צריך לקשר שורה או להשאיר את השורות נפרדות."}
    if action == "split":
        if not item["capture_id"]:
            return {"status": item["status"], "missing": "אין קליטה לפיצול."}
        result = split_capture(conn, item["capture_id"], body.get("new_capture_id") or str(uuid.uuid4()), body.get("photo_ids") or [], created_at)
        return {"status": item["status"], "split": result}
    raise ValueError("פעולה לא מוכרת.")


def row_documentation(conn: sqlite3.Connection) -> dict[str, dict]:
    documented: set[str] = set()
    for photo in conn.execute("SELECT shows_json FROM photos"):
        shows = parse_json(photo["shows_json"], {})
        if shows.get("scope") == "rows":
            documented.update(shows.get("row_ids") or [])
    pending: set[str] = set()
    identity: dict[str, str] = {}
    for link in conn.execute("SELECT inventory_row_id, review_status FROM evidence_links"):
        row_id = link["inventory_row_id"]
        status = link["review_status"]
        if status == "pending_review":
            pending.add(row_id)
            identity[row_id] = "pending"
        elif status == "user_confirmed" and identity.get(row_id) != "pending":
            identity[row_id] = "user_confirmed"
        elif status == "auto_linked" and row_id not in identity:
            identity[row_id] = "auto_linked"
    review_rows: set[str] = set()
    for review in conn.execute(
        "SELECT row_ids_json, status FROM review_items WHERE status IN ('open', 'deferred') AND queue != 'valuation'"
    ):
        review_rows.update(parse_json(review["row_ids_json"], []))
    flags = {}
    for row in conn.execute("SELECT id, kind, listed_area, sheet_name, duplicate_group FROM inventory_rows"):
        if row["kind"] != "equipment":
            continue
        row_id = row["id"]
        flags[row_id] = {
            "documented": row_id in documented,
            "needs_review": row_id in pending or row_id in review_rows,
            "identity": identity.get(row_id, "none"),
            "listed_area": row["listed_area"] or "",
            "sheet_name": row["sheet_name"],
            "duplicate_group": row["duplicate_group"] or "",
        }
    return flags


def visit_summary(conn: sqlite3.Connection) -> dict:
    flags = row_documentation(conn)
    by_area: dict[str, dict] = {}
    for row_id, flag in flags.items():
        area = flag["listed_area"] or "לא ידוע"
        bucket = by_area.setdefault(area, {"area": area, "rows": 0, "documented": 0, "needs_review": 0})
        bucket["rows"] += 1
        if flag["documented"]:
            bucket["documented"] += 1
        if flag["needs_review"]:
            bucket["needs_review"] += 1
    documented = sum(1 for flag in flags.values() if flag["documented"])
    needs_review = sum(1 for flag in flags.values() if flag["needs_review"])
    equipment = len(flags)
    captures = [dict(row) for row in conn.execute("SELECT * FROM captures")]
    synced = sum(1 for capture in captures if capture["sync_status"] == "synced")
    pending_sync = sum(1 for capture in captures if capture["sync_status"] != "synced")
    unresolved = 0
    for capture in captures:
        if not capture["finalized"]:
            continue
        links = capture_links(conn, capture["id"])
        confirmed = [link for link in links if link["review_status"] in {"auto_linked", "user_confirmed"}]
        if not confirmed:
            unresolved += 1
    revisit = []
    for row in conn.execute(
        """
        SELECT id, tag_original, description, listed_area, sheet_name, original_row
        FROM inventory_rows WHERE kind = 'equipment'
        """
    ):
        flag = flags[row["id"]]
        if flag["documented"] and not flag["needs_review"]:
            continue
        if flag["documented"] and flag["needs_review"]:
            reason = "יש תיעוד, ועדיין יש שאלה פתוחה"
        elif flag["needs_review"]:
            reason = "שאלה פתוחה"
        else:
            reason = "אין עדיין צילום של הפריט עצמו"
        if not flag["documented"] or flag["needs_review"]:
            revisit.append(
                {
                    "row_id": row["id"],
                    "tag": row["tag_original"] or "",
                    "description": row["description"] or "",
                    "listed_area": row["listed_area"] or "לא ידוע",
                    "sheet_name": row["sheet_name"],
                    "original_row": row["original_row"],
                    "reason": reason,
                    "documented": flag["documented"],
                    "needs_review": flag["needs_review"],
                }
            )
    visit = conn.execute("SELECT * FROM visits WHERE is_current = 1").fetchone()
    return {
        "visit": dict(visit) if visit else None,
        "captures_synced": synced,
        "captures_not_synced": pending_sync,
        "rows_equipment": equipment,
        "rows_documented": documented,
        "rows_needs_review": needs_review,
        "rows_not_documented": equipment - documented,
        "unresolved_captures": unresolved,
        "by_area": sorted(by_area.values(), key=lambda item: item["area"]),
        "revisit_count": len(revisit),
        "footnote": "המספרים סופרים שורות מלאי, לא מכונות פיזיות. תגים כפולים וכמויות שלא הוכרעו נשארים שורות נפרדות.",
        "revisit": revisit,
    }


def public_row(record: sqlite3.Row, source_name: str) -> dict:
    return {
        "id": record["id"],
        "kind": record["kind"],
        "source_file": source_name,
        "source_kind": record["source_kind"],
        "sheet_name": record["sheet_name"],
        "original_row": record["original_row"],
        "identity_area": record["identity_area"],
        "listed_area": record["listed_area"] or "",
        "tag_original": record["tag_original"] or "",
        "tag_norm": record["tag_norm"] or "",
        "description": record["description"] or "",
        "manufacturer": record["manufacturer"] or "",
        "model": record["model"] or "",
        "serial": record["serial"] or "",
        "serial_norm": record["serial_norm"] or "",
        "quantity": record["quantity"] or "",
        "remarks": record["remarks"] or "",
        "specs": record["specs"] or "",
        "material": record["material"] or "",
        "system_number": record["system_number"] or "",
        "efd": record["efd"] or "",
        "installation_raw": record["installation_raw"] or "",
        "labels": parse_json(record["labels_json"], []),
        "duplicate_group": record["duplicate_group"] or "",
        "prefix_category": __import__("app.machines", fromlist=["prefix_label"]).prefix_label(record["tag_original"] or record["tag_norm"] or ""),
        "raw": parse_json(record["raw_json"], {}),
    }


def inventory_csv(conn: sqlite3.Connection) -> str:
    flags = row_documentation(conn)
    observed: dict[str, list[str]] = {}
    notes: dict[str, list[str]] = {}
    photos: dict[str, list[str]] = {}
    for link in conn.execute(
        """
        SELECT l.inventory_row_id, c.observed_area_name, c.note, c.id AS capture_id
        FROM evidence_links l JOIN captures c ON c.id = l.capture_id
        """
    ):
        if link["observed_area_name"]:
            observed.setdefault(link["inventory_row_id"], []).append(link["observed_area_name"])
        if link["note"]:
            notes.setdefault(link["inventory_row_id"], []).append(link["note"])
    for photo in conn.execute("SELECT id, shows_json, capture_id FROM photos"):
        shows = parse_json(photo["shows_json"], {})
        if shows.get("scope") != "rows":
            continue
        for row_id in shows.get("row_ids") or []:
            photos.setdefault(row_id, []).append(f"{photo['capture_id']}/{photo['id']}")
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "stable_id",
            "source_file",
            "sheet_or_table",
            "original_row",
            "tag",
            "description",
            "listed_area",
            "observed_areas",
            "notes",
            "documentation_status",
            "identity_status",
            "labels",
            "quantity",
            "duplicate_group",
            "photo_references",
        ]
    )
    sources = {row["id"]: row["display_name"] for row in conn.execute("SELECT id, display_name FROM source_files")}
    for record in conn.execute("SELECT * FROM inventory_rows WHERE kind = 'equipment' ORDER BY sheet_name, original_row"):
        flag = flags.get(record["id"], {})
        if flag.get("needs_review"):
            doc = "needs_review"
        elif flag.get("documented"):
            doc = "documented"
        else:
            doc = "not_documented"
        writer.writerow(
            [
                record["id"],
                sources.get(record["source_file_id"], ""),
                record["sheet_name"],
                record["original_row"],
                record["tag_original"] or "",
                record["description"] or "",
                record["listed_area"] or "",
                " | ".join(dict.fromkeys(observed.get(record["id"], []))),
                " | ".join(notes.get(record["id"], [])),
                doc,
                flag.get("identity", "none"),
                " | ".join(parse_json(record["labels_json"], [])),
                record["quantity"] or "",
                record["duplicate_group"] or "",
                " | ".join(photos.get(record["id"], [])),
            ]
        )
    return buffer.getvalue()


def export_package(conn: sqlite3.Connection, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    flags = row_documentation(conn)
    manifest = {
        "exported_at": now_iso(),
        "note": "קישור תמונה לשורת מלאי אינו קשר מערכת ואינו איחוד של פריטים.",
        "captures": [],
        "links": [],
    }
    captures = []
    for capture in conn.execute("SELECT * FROM captures"):
        photos = []
        for photo in conn.execute("SELECT * FROM photos WHERE capture_id = ?", (capture["id"],)):
            shows = parse_json(photo["shows_json"], {})
            photos.append(
                {
                    "photo_id": photo["id"],
                    "role": photo["role"],
                    "file": f"photos/{capture['id']}/{photo['id']}{Path(photo['stored_path']).suffix if photo['stored_path'] else ''}",
                    "shows": shows,
                    "upload_state": photo["upload_state"],
                }
            )
        captures.append(
            {
                "capture_id": capture["id"],
                "observed_area": capture["observed_area_name"],
                "note": capture["note"],
                "info_provided_by": capture["info_provided_by"],
                "tag_text": capture["tag_text"],
                "sync_status": capture["sync_status"],
                "photos": photos,
            }
        )
    manifest["captures"] = captures
    for link in conn.execute(
        """
        SELECT l.*, r.sheet_name, r.original_row, r.tag_original, r.id AS row_id
        FROM evidence_links l JOIN inventory_rows r ON r.id = l.inventory_row_id
        """
    ):
        manifest["links"].append(
            {
                "capture_id": link["capture_id"],
                "inventory_row_id": link["row_id"],
                "sheet_or_table": link["sheet_name"],
                "original_row": link["original_row"],
                "tag": link["tag_original"],
                "method": link["method"],
                "review_status": link["review_status"],
                "reason": link["reason"],
                "documented": flags.get(link["row_id"], {}).get("documented", False),
            }
        )
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("inventory.csv", inventory_csv(conn))
        for capture in captures:
            for photo in conn.execute("SELECT * FROM photos WHERE capture_id = ?", (capture["capture_id"],)):
                if not photo["stored_path"] or not Path(photo["stored_path"]).exists():
                    continue
                suffix = Path(photo["stored_path"]).suffix
                archive.write(photo["stored_path"], f"photos/{capture['capture_id']}/{photo['id']}{suffix}")
    return destination
