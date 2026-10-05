"""Load the plant workbook and the laboratory document.

One plant file is the master inventory. The laboratory list is in scope and is
stored as its own source rows. Generic identical names are not collapsed.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from docx import Document

from app.matching import norm_id

LABEL_PATTERNS = [
    ("NEW PROJECT", re.compile(r"\bNEW PROJECT\b", re.I)),
    ("UNDER MAINTENANCE", re.compile(r"\bUNDER MAINTENANCE\b", re.I)),
    ("NOT IN USE", re.compile(r"\bNOT IN USE\b", re.I)),
    ("FUTURE", re.compile(r"\bFUTURE\b", re.I)),
    ("HOLD", re.compile(r"\bHOLD\b", re.I)),
    ("SPARE", re.compile(r"\bSPARE\b", re.I)),
]
DATE_RE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
RANGE_TAG_RE = re.compile(r"^\d+\s*-\s*\d+$")
SERIAL_TOKEN_RE = re.compile(r"[A-Za-z0-9-]{6,}")
GENERIC_NAMES = {"centrifuge": "Centrifuge", "dosing pump": "Dosing Pump"}

PLANT_MAPPING = {
    "area_source": "sheet",
    "columns": {
        "tag": "Item",
        "description": "Description",
        "remarks": "Remarks",
        "system_number": "System",
        "efd": "Efd",
        "specs": "Character",
        "material": "Material",
    },
}
LAB_MAPPING = {
    "area_source": "unknown",
    "columns": {
        "description": "Instrument Name",
        "tag": "Instrument Tag No.",
        "manufacturer": "Manufacturer",
        "model": "Model No.",
        "serial": "Serial No.",
        "installation": "Date of Installation",
    },
}


def file_sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_row_id(file_sha256: str, sheet: str, row_number: int) -> str:
    digest = hashlib.sha256(f"{file_sha256}|{sheet}|{row_number}".encode()).hexdigest()
    return "ROW-" + digest[:16]


def area_id(name: str) -> str:
    digest = hashlib.sha256(name.strip().casefold().encode()).hexdigest()
    return "AREA-" + digest[:12]


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def extract_labels(*texts: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for text in texts:
        if not text:
            continue
        for canon, pattern in LABEL_PATTERNS:
            match = pattern.search(str(text))
            if match and canon not in seen:
                seen.add(canon)
                found.append(match.group(0))
    return found


def classify_date(value: str) -> str:
    text = (value or "").strip()
    match = DATE_RE.fullmatch(text)
    if not match:
        return "not_date"
    day, month, year = (int(part) for part in match.groups())
    try:
        date(year, month, day)
    except ValueError:
        return "invalid"
    return "valid"


def two_serials(value: str) -> bool:
    parts = (value or "").split()
    if len(parts) != 2:
        return False
    return all(SERIAL_TOKEN_RE.fullmatch(part) for part in parts)


def filled_signature(raw: dict) -> str:
    parts = []
    for key in sorted(raw):
        text = cell_text(raw.get(key))
        if text:
            parts.append(f"{key}={text}")
    return "|".join(parts)


def _mapped(raw: dict, mapping: dict, field: str) -> str:
    header = (mapping.get("columns") or {}).get(field)
    if not header:
        return ""
    return cell_text(raw.get(header))


def listed_area_for(raw: dict, mapping: dict, sheet_name: str) -> str:
    source = mapping.get("area_source") or "sheet"
    if source == "sheet":
        return sheet_name
    if source == "unknown":
        return ""
    if source.startswith("column:"):
        return cell_text(raw.get(source.split(":", 1)[1]))
    return ""


def mapping_uses_price(mapping: dict) -> str | None:
    for field, header in (mapping.get("columns") or {}).items():
        if header and str(header).strip().casefold() == "price":
            return field
    area = mapping.get("area_source") or ""
    if area.lower().endswith("price") or area.casefold() == "column:price":
        return "area"
    return None


def read_plant_xlsx(path: Path) -> tuple[list[dict], list[str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    records: list[dict] = []
    sheets: list[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheets.append(sheet_name)
            worksheet = workbook[sheet_name]
            header_map: dict[int, str] = {}
            for index, row in enumerate(worksheet.iter_rows(max_col=40, values_only=True), start=1):
                if index < 4:
                    continue
                if index == 4:
                    for col, value in enumerate(row):
                        text = cell_text(value)
                        if text:
                            header_map[col] = text
                    continue
                if not header_map:
                    continue
                raw = {header_map[col]: cell_text(row[col]) for col in header_map if col < len(row)}
                if not any(raw.values()):
                    continue
                records.append({"sheet_name": sheet_name, "original_row": index, "raw": raw, "source_kind": "plant"})
    finally:
        workbook.close()
    return records, sheets


def read_plant_csv(path: Path) -> tuple[list[dict], list[str]]:
    text = path.read_text(encoding="utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    records = []
    for index, row in enumerate(reader, start=2):
        raw = {key: cell_text(value) for key, value in row.items() if key}
        if not any(raw.values()):
            continue
        records.append({"sheet_name": "csv", "original_row": index, "raw": raw, "source_kind": "plant"})
    return records, ["csv"]


def _lab_cells(row) -> list[str]:
    return [cell_text(cell.text.replace("\n", " ")) for cell in row.cells]


def is_lab_banner(cells: list[str]) -> bool:
    if not cells:
        return True
    if cells[0] == "Instrument Name":
        return True
    if cells[0] == "SRD List":
        return True
    if all(not cell or cell == "SRD List" for cell in cells):
        return True
    return False


LAB_HEADER_FIELDS = {
    "instrument name": "description",
    "name": "description",
    "instrument tag no.": "tag",
    "instrument tag": "tag",
    "tag no.": "tag",
    "tag": "tag",
    "manufacturer": "manufacturer",
    "model no.": "model",
    "model": "model",
    "model number": "model",
    "serial no.": "serial",
    "serial": "serial",
    "serial number": "serial",
    "date of installation": "installation",
    "installation": "installation",
    "install date": "installation",
}


def header_key(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().casefold())


def lab_mapping_for(headers: list[str]) -> dict | None:
    columns: dict[str, str] = {}
    for header in headers:
        field = LAB_HEADER_FIELDS.get(header_key(header))
        if field and field not in columns and header.strip():
            columns[field] = header.strip()
    if not {"tag", "manufacturer", "model"} <= set(columns):
        return None
    return {"area_source": "unknown", "columns": columns}


def _is_lab_header_cells(cells: list[str]) -> bool:
    return lab_mapping_for(cells) is not None


def read_lab_docx(path: Path) -> list[dict]:
    document = Document(str(path))
    headers = [
        "Instrument Name",
        "Instrument Tag No.",
        "Manufacturer",
        "Model No.",
        "Serial No.",
        "Date of Installation",
    ]
    records = []
    for table_index, table in enumerate(document.tables, start=1):
        sheet_name = f"טבלה {table_index}"
        for row_index, row in enumerate(table.rows, start=1):
            cells = _lab_cells(row)
            if row_index == 1:
                continue
            while len(cells) < 6:
                cells.append("")
            raw = dict(zip(headers, cells[:6]))
            if not any(raw.values()):
                continue
            records.append(
                {
                    "sheet_name": sheet_name,
                    "original_row": row_index,
                    "raw": raw,
                    "source_kind": "lab",
                    "banner": is_lab_banner(cells[:6]),
                }
            )
    return records


def _logical_fields(raw: dict, mapping: dict, sheet_name: str, source_kind: str) -> dict:
    tag_original = _mapped(raw, mapping, "tag")
    description = _mapped(raw, mapping, "description")
    serial = _mapped(raw, mapping, "serial")
    installation = _mapped(raw, mapping, "installation")
    texts = list(raw.values())
    labels = extract_labels(*texts)
    installation_raw = installation
    if installation and any(label.casefold() == installation.casefold() for label in labels):
        installation_raw = ""
    if classify_date(installation) == "invalid":
        installation_raw = installation
    serial_norm = "" if two_serials(serial) else norm_id(serial)
    identity = sheet_name if source_kind == "plant" else f"lab:{sheet_name}"
    kind = "equipment"
    if source_kind == "plant" and not tag_original and not description:
        kind = "not_equipment"
    if source_kind == "lab" and raw.get("_banner"):
        kind = "not_equipment"
    return {
        "kind": kind,
        "identity_area": identity,
        "listed_area": listed_area_for(raw, mapping, sheet_name),
        "tag_original": tag_original,
        "tag_norm": norm_id(tag_original),
        "description": description,
        "manufacturer": _mapped(raw, mapping, "manufacturer"),
        "model": _mapped(raw, mapping, "model"),
        "serial": serial,
        "serial_norm": serial_norm,
        "quantity": _mapped(raw, mapping, "quantity"),
        "remarks": _mapped(raw, mapping, "remarks"),
        "specs": _mapped(raw, mapping, "specs"),
        "material": _mapped(raw, mapping, "material"),
        "system_number": _mapped(raw, mapping, "system_number"),
        "efd": _mapped(raw, mapping, "efd"),
        "installation_raw": installation_raw,
        "labels": labels,
    }


def assign_duplicate_groups(rows: list[dict]) -> list[dict]:
    """Return review payloads for tags that are not an identical same-area duplicate."""
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        if row["kind"] != "equipment" or not row["tag_norm"]:
            continue
        grouped[row["tag_norm"]].append(row)
    for row in rows:
        row["duplicate_group"] = None
    reviews = []
    for tag, group in grouped.items():
        if len(group) < 2:
            continue
        by_area: dict[str, list[dict]] = defaultdict(list)
        for row in group:
            by_area[row["identity_area"]].append(row)
        conflict = len(by_area) > 1
        for area, area_rows in by_area.items():
            if len(area_rows) < 2:
                continue
            signatures = {filled_signature(row["raw"]) for row in area_rows}
            if len(signatures) == 1:
                signature = next(iter(signatures))
                group_id = "DUP-" + hashlib.sha256(f"{area}|{tag}|{signature}".encode()).hexdigest()[:12]
                for row in area_rows:
                    row["duplicate_group"] = group_id
            else:
                conflict = True
        if conflict:
            reviews.append(
                {
                    "kind": "duplicate_tag",
                    "queue": "import",
                    "dedupe_key": f"import:dup-tag:{tag}",
                    "priority": "high",
                    "question": (
                        f"התג {tag} מופיע יותר מפעם אחת עם נתונים שונים או באזורים שונים. "
                        "השורות נשארו נפרדות. איזו שורה מתארת את הפריט, או שאלה באמת פריטים שונים?"
                    ),
                    "row_ids": [row["id"] for row in group],
                    "payload": {"tag": tag, "places": _places(group)},
                }
            )
    return reviews


def _places(rows: list[dict]) -> list[dict]:
    places = []
    for row in rows:
        places.append(
            {
                "row_id": row["id"],
                "sheet_name": row["sheet_name"],
                "original_row": row["original_row"],
                "description": row["description"],
                "listed_area": row["listed_area"],
                "tag": row["tag_original"],
            }
        )
    return places


def generic_name_reviews(rows: list[dict]) -> list[dict]:
    reviews = []
    for key, label in GENERIC_NAMES.items():
        matched = []
        for row in rows:
            if row["kind"] != "equipment":
                continue
            description = (row["description"] or "").strip().casefold()
            if description == key:
                matched.append(row)
        if len(matched) < 2:
            continue
        plant = [row for row in matched if row["source_kind"] == "plant"]
        lab = [row for row in matched if row["source_kind"] == "lab"]
        reviews.append(
            {
                "kind": "generic_name",
                "queue": "import",
                "dedupe_key": f"import:generic:{key}",
                "priority": "medium",
                "question": (
                    f"השם {label} זהה בכמה שורות. לא אוחדו לפריט אחד: "
                    "השם כללי, ולכל שורה תג אחר. "
                    f"במפעל {len(plant)} שורות ובמעבדה {len(lab)}. "
                    "אין כאן אימות שזה אותו פריט פיזי."
                ),
                "row_ids": [row["id"] for row in matched],
                "payload": {"name": label, "places": _places(matched)},
            }
        )
    return reviews


def lab_anomaly_reviews(rows: list[dict]) -> list[dict]:
    reviews = []
    panels = []
    sensors = []
    bad_dates = []
    multi_serial = []
    for row in rows:
        if row["kind"] != "equipment" or row["source_kind"] != "lab":
            continue
        tag = (row["tag_original"] or "").strip()
        description = row["description"] or ""
        raw_date = row["raw"].get("Date of Installation", "")
        if "panel" in description.casefold() or RANGE_TAG_RE.fullmatch(tag):
            panels.append(row)
        if "sensor" in description.casefold() and norm_id(tag) == "":
            sensors.append(row)
        if classify_date(raw_date) == "invalid":
            bad_dates.append(row)
        if two_serials(row["serial"] or ""):
            multi_serial.append(row)
    if panels:
        reviews.append(
            {
                "kind": "lab_anomaly",
                "queue": "import",
                "dedupe_key": "import:lab:panels",
                "priority": "medium",
                "question": (
                    "חלק משורות המעבדה אינן מכונה אחת: התג הוא טווח (למשל 1-60) או שם של פאנל גז. "
                    "כמה יחידות פיזיות יש בכל שורה? שימו לב שהמספר 63 מופיע גם בטווח 61-63 וגם ב-63-66."
                ),
                "row_ids": [row["id"] for row in panels],
                "payload": {"places": _places(panels)},
            }
        )
    if sensors:
        reviews.append(
            {
                "kind": "lab_anomaly",
                "queue": "import",
                "dedupe_key": "import:lab:sensors",
                "priority": "medium",
                "question": (
                    "יש חיישנים בלי תג זיהוי אמיתי (התא ריק, '-' או '--'). "
                    "שניים מהם הם חיישני חמצן. איך מבדילים ביניהם בסיור?"
                ),
                "row_ids": [row["id"] for row in sensors],
                "payload": {"places": _places(sensors)},
            }
        )
    if bad_dates:
        shown = ", ".join(sorted({row["raw"].get("Date of Installation", "") for row in bad_dates}))
        reviews.append(
            {
                "kind": "lab_anomaly",
                "queue": "import",
                "dedupe_key": "import:lab:bad-date",
                "priority": "medium",
                "question": (
                    f"בתאריך ההתקנה כתוב {shown}. אין תאריך כזה בלוח השנה. "
                    "מה התאריך הנכון? הערך המקורי נשמר כמו שהוא."
                ),
                "row_ids": [row["id"] for row in bad_dates],
                "payload": {"places": _places(bad_dates)},
            }
        )
    if multi_serial:
        reviews.append(
            {
                "kind": "lab_anomaly",
                "queue": "import",
                "dedupe_key": "import:lab:two-serials",
                "priority": "medium",
                "question": (
                    "בשורה של Auto Titrator כתובים שני מספרים סידוריים באותו תא. "
                    "האם זה מכשיר אחד? לא נבחר אחד מהם ולא הומצא מספר."
                ),
                "row_ids": [row["id"] for row in multi_serial],
                "payload": {"places": _places(multi_serial), "serials": [row["serial"] for row in multi_serial]},
            }
        )
    return reviews


def _build_rows(records: list[dict], mapping: dict, file_sha256: str, source_kind: str) -> list[dict]:
    rows = []
    for record in records:
        raw = dict(record["raw"])
        if record.get("banner"):
            raw["_banner"] = "1"
        logical = _logical_fields(raw, mapping, record["sheet_name"], source_kind)
        if record.get("banner"):
            logical["kind"] = "not_equipment"
            raw.pop("_banner", None)
        raw.pop("_banner", None)
        rows.append(
            {
                "id": stable_row_id(file_sha256, record["sheet_name"], record["original_row"]),
                "sheet_name": record["sheet_name"],
                "original_row": record["original_row"],
                "raw": raw,
                "source_kind": source_kind,
                "duplicate_group": None,
                **logical,
            }
        )
    return rows


def replace_file_rows(conn: sqlite3.Connection, file_id: str, rows: list[dict]) -> None:
    conn.execute("DELETE FROM inventory_rows WHERE source_file_id = ?", (file_id,))
    conn.executemany(
        """
        INSERT INTO inventory_rows (
          id, source_file_id, kind, sheet_name, original_row, identity_area, listed_area,
          tag_original, tag_norm, description, manufacturer, model, serial, serial_norm,
          quantity, remarks, specs, material, system_number, efd, installation_raw,
          labels_json, raw_json, duplicate_group, source_kind
        ) VALUES (
          :id, :source_file_id, :kind, :sheet_name, :original_row, :identity_area, :listed_area,
          :tag_original, :tag_norm, :description, :manufacturer, :model, :serial, :serial_norm,
          :quantity, :remarks, :specs, :material, :system_number, :efd, :installation_raw,
          :labels_json, :raw_json, :duplicate_group, :source_kind
        )
        """,
        [
            {
                "source_file_id": file_id,
                "labels_json": json.dumps(row["labels"], ensure_ascii=False),
                "raw_json": json.dumps(row["raw"], ensure_ascii=False),
                **{key: row[key] for key in row if key not in {"labels", "raw"}},
            }
            for row in rows
        ],
    )


def upsert_reviews(conn: sqlite3.Connection, reviews: list[dict], now: str) -> None:
    for review in reviews:
        existing = conn.execute(
            "SELECT id, status FROM review_items WHERE dedupe_key = ?",
            (review["dedupe_key"],),
        ).fetchone()
        payload = json.dumps(review.get("payload") or {}, ensure_ascii=False)
        row_ids = json.dumps(review["row_ids"], ensure_ascii=False)
        if existing:
            if existing["status"] == "open":
                conn.execute(
                    """
                    UPDATE review_items
                    SET question = ?, row_ids_json = ?, payload_json = ?, updated_at = ?, priority = ?
                    WHERE id = ?
                    """,
                    (review["question"], row_ids, payload, now, review["priority"], existing["id"]),
                )
            continue
        review_id = "REV-" + hashlib.sha256(review["dedupe_key"].encode()).hexdigest()[:16]
        conn.execute(
            """
            INSERT INTO review_items (
              id, kind, queue, dedupe_key, question, priority, status, capture_id,
              row_ids_json, payload_json, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, 'open', NULL, ?, ?, ?, ?)
            """,
            (
                review_id,
                review["kind"],
                review["queue"],
                review["dedupe_key"],
                review["question"],
                review["priority"],
                row_ids,
                payload,
                now,
                now,
            ),
        )


def refresh_import_reviews(conn: sqlite3.Connection, now: str) -> None:
    rows = []
    for record in conn.execute("SELECT * FROM inventory_rows"):
        item = dict(record)
        item["raw"] = json.loads(item["raw_json"])
        item["labels"] = json.loads(item["labels_json"])
        rows.append(item)
    reviews = []
    reviews.extend(assign_duplicate_groups(rows))
    conn.executemany(
        "UPDATE inventory_rows SET duplicate_group = ? WHERE id = ?",
        [(row.get("duplicate_group"), row["id"]) for row in rows],
    )
    reviews.extend(generic_name_reviews(rows))
    reviews.extend(lab_anomaly_reviews(rows))
    upsert_reviews(conn, reviews, now)
    keep = {review["dedupe_key"] for review in reviews}
    for item in conn.execute("SELECT id, dedupe_key FROM review_items WHERE queue = 'import' AND status = 'open'"):
        if item["dedupe_key"] not in keep:
            conn.execute("DELETE FROM review_items WHERE id = ?", (item["id"],))


def ensure_areas(conn: sqlite3.Connection, sheet_names: list[str]) -> None:
    for name in sheet_names:
        conn.execute(
            """
            INSERT OR IGNORE INTO areas (id, name, parent_id, source, note)
            VALUES (?, ?, NULL, 'file', 'שם הגיליון ברשימת המפעל. זו יחידת האזור שבקובץ, לא כתובת.')
            """,
            (area_id(name), name),
        )
    conn.execute(
        """
        INSERT OR IGNORE INTO areas (id, name, parent_id, source, note)
        VALUES (?, 'מעבדה', NULL, 'list', 'רשימת המעבדה כלולה בסיור. הקובץ לא מציין חדר.')
        """,
        (area_id("מעבדה"),),
    )
    conn.execute(
        """
        INSERT OR IGNORE INTO areas (id, name, parent_id, source, note)
        VALUES (?, 'לא ידוע', NULL, 'list', 'אזור לא ידוע הוא ערך תקין.')
        """,
        (area_id("לא ידוע"),),
    )


def write_import_notes(conn: sqlite3.Connection) -> None:
    conn.execute("DELETE FROM import_notes")
    label_counts: dict[str, int] = defaultdict(int)
    equipment = 0
    for record in conn.execute("SELECT kind, labels_json, source_kind FROM inventory_rows"):
        if record["kind"] != "equipment":
            continue
        equipment += 1
        for label in json.loads(record["labels_json"]):
            label_counts[label.upper() if label.isascii() else label] += 1
    notes = [
        "נטען קובץ מפעל אחד. עותק שני של אותה רשימה לא נטען, כדי לא לספור את המלאי פעמיים.",
        "רשימת מכשירי המעבדה כלולה בסיור ונשמרת כשורות מקור נפרדות.",
        "העמודה price באקסל אינה מחיר מכירה ולא מופתה לשדה. אין בקובץ מחירי מכירה.",
        "ההדגשה הירוקה במסמך המעבדה (Perrigo) אינה שאלת סיור ולא נוצרה עבורה משימה.",
        "שמות כלליים זהים (Centrifuge, Dosing Pump) לא אוחדו. השאלה פתוחה בתוך בדיקות הקובץ.",
        "שורות FUTURE, HOLD, SPARE, NEW PROJECT, Not in Use ו-Under maintenance נשארו במלאי עם הסימון המקורי.",
        "קישור צילום לשורה אינו קשר מערכת, ואינו אומר שהפריטים נמכרים יחד.",
        f"נספרו {equipment} שורות ציוד. זה מספר שורות, לא מספר מכונות פיזיות.",
    ]
    if label_counts:
        pretty = ", ".join(f"{name} ({count})" for name, count in sorted(label_counts.items()))
        notes.append(f"סימונים שנשמרו כמו בקובץ: {pretty}.")
    conn.executemany("INSERT INTO import_notes (level, text) VALUES ('info', ?)", [(note,) for note in notes])


def ensure_current_visit(conn: sqlite3.Connection, now: str) -> None:
    current = conn.execute("SELECT id FROM visits WHERE is_current = 1").fetchone()
    if current:
        return
    visit_id = "VISIT-" + hashlib.sha256(now.encode()).hexdigest()[:12]
    conn.execute(
        """
        INSERT INTO visits (id, visit_date, description, is_current, created_at)
        VALUES (?, ?, 'סיור ציוד', 1, ?)
        """,
        (visit_id, now[:10], now),
    )


def import_parsed_file(
    conn: sqlite3.Connection,
    *,
    kind: str,
    filename: str,
    display_name: str,
    sha256: str,
    records: list[dict],
    sheet_names: list[str],
    mapping: dict,
    now: str,
    replace: bool,
) -> dict:
    blocked = mapping_uses_price(mapping)
    if blocked:
        raise ValueError("העמודה price אינה מחיר מכירה ולא ניתן למפות אותה לשדה מלאי.")
    existing = conn.execute("SELECT id, sha256, kind FROM source_files WHERE kind = ?", (kind,)).fetchone()
    if existing and existing["sha256"] != sha256 and not replace:
        raise FileExistsError("כבר טעון קובץ מפעל. טעינה נוספת בלי החלפה תספור את המלאי פעמיים." if kind == "plant" else "רשימת המעבדה כבר טעונה.")
    if existing and existing["sha256"] == sha256 and not replace:
        return {"status": "already_loaded", "file_id": existing["id"], "rows": 0}
    file_id = "FILE-" + sha256[:12]
    if existing and existing["id"] != file_id:
        conn.execute("DELETE FROM inventory_rows WHERE source_file_id = ?", (existing["id"],))
        conn.execute("DELETE FROM source_files WHERE id = ?", (existing["id"],))
    conn.execute(
        """
        INSERT INTO source_files (id, kind, filename, display_name, sha256, imported_at, mapping_json)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(id) DO UPDATE SET
          filename = excluded.filename,
          display_name = excluded.display_name,
          mapping_json = excluded.mapping_json,
          imported_at = excluded.imported_at
        """,
        (file_id, kind, filename, display_name, sha256, now, json.dumps(mapping, ensure_ascii=False)),
    )
    rows = _build_rows(records, mapping, sha256, kind)
    replace_file_rows(conn, file_id, rows)
    if kind == "plant":
        ensure_areas(conn, sheet_names)
    return {"status": "imported", "file_id": file_id, "rows": len(rows)}


def seed_sources(conn: sqlite3.Connection, sources_dir: Path, now: str) -> None:
    plant = sources_dir / "plant-equipment.xlsx"
    lab = sources_dir / "lab-instruments.docx"
    if not plant.exists() or not lab.exists():
        return
    loaded = conn.execute("SELECT COUNT(*) AS n FROM source_files").fetchone()["n"]
    if loaded:
        return
    plant_records, sheets = read_plant_xlsx(plant)
    import_parsed_file(
        conn,
        kind="plant",
        filename=plant.name,
        display_name="רשימת ציוד המפעל",
        sha256=file_sha(plant),
        records=plant_records,
        sheet_names=sheets,
        mapping=PLANT_MAPPING,
        now=now,
        replace=False,
    )
    lab_records = read_lab_docx(lab)
    import_parsed_file(
        conn,
        kind="lab",
        filename=lab.name,
        display_name="רשימת מכשירי מעבדה",
        sha256=file_sha(lab),
        records=lab_records,
        sheet_names=[],
        mapping=LAB_MAPPING,
        now=now,
        replace=False,
    )
    ensure_areas(conn, sheets)
    refresh_import_reviews(conn, now)
    write_import_notes(conn)
    ensure_current_visit(conn, now)


def remap_file(conn: sqlite3.Connection, kind: str, mapping: dict, now: str) -> dict:
    blocked = mapping_uses_price(mapping)
    if blocked:
        raise ValueError("העמודה price אינה מחיר מכירה ולא ניתן למפות אותה לשדה מלאי.")
    source = conn.execute("SELECT * FROM source_files WHERE kind = ?", (kind,)).fetchone()
    if not source:
        raise LookupError("הקובץ לא נמצא.")
    rows = conn.execute(
        "SELECT * FROM inventory_rows WHERE source_file_id = ?",
        (source["id"],),
    ).fetchall()
    records = []
    for row in rows:
        raw = json.loads(row["raw_json"])
        records.append(
            {
                "sheet_name": row["sheet_name"],
                "original_row": row["original_row"],
                "raw": raw,
                "source_kind": kind,
                "banner": row["kind"] == "not_equipment" and kind == "lab" and raw.get("Instrument Name") in {"Instrument Name", "SRD List"},
            }
        )
    # Recompute from raw with the new mapping. Stable IDs stay on sheet+row.
    built = _build_rows(records, mapping, source["sha256"], kind)
    for record, built_row in zip(records, built):
        if record.get("banner"):
            built_row["kind"] = "not_equipment"
    replace_file_rows(conn, source["id"], built)
    conn.execute(
        "UPDATE source_files SET mapping_json = ?, imported_at = ? WHERE id = ?",
        (json.dumps(mapping, ensure_ascii=False), now, source["id"]),
    )
    refresh_import_reviews(conn, now)
    write_import_notes(conn)
    return {"status": "remapped", "rows": len(built)}


def preview_workbook(path: Path) -> dict:
    records, sheets = read_plant_xlsx(path) if path.suffix.lower() == ".xlsx" else read_plant_csv(path)
    by_sheet: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        by_sheet[record["sheet_name"]].append(record)
    sheet_info = []
    for name in sheets:
        sample_rows = by_sheet.get(name, [])[:2]
        headers = list(sample_rows[0]["raw"].keys()) if sample_rows else []
        sheet_info.append(
            {
                "name": name,
                "rows": len(by_sheet.get(name, [])),
                "headers": headers,
                "samples": [{"row": row["original_row"], "cells": row["raw"]} for row in sample_rows],
            }
        )
    return {"sheets": sheet_info, "suggested_mapping": PLANT_MAPPING}


def _records_from_header_rows(sheet_name: str, header_row: int, headers: list[str], data_rows: list[tuple[int, list[str]]], source_kind: str) -> list[dict]:
    records = []
    for row_number, cells in data_rows:
        raw = {headers[index]: cells[index] if index < len(cells) else "" for index in range(len(headers)) if headers[index].strip()}
        if not any(str(value).strip() for value in raw.values()):
            continue
        banner = _is_lab_header_cells([raw.get(header, "") for header in headers]) if source_kind == "lab" else False
        records.append(
            {
                "sheet_name": sheet_name,
                "original_row": row_number,
                "raw": raw,
                "source_kind": source_kind,
                "banner": banner or (source_kind == "lab" and is_lab_banner([raw.get(header, "") for header in headers[:6]])),
            }
        )
    return records


def parse_lab_docx(path: Path) -> dict | None:
    document = Document(str(path))
    records: list[dict] = []
    mapping = None
    for table_index, table in enumerate(document.tables, start=1):
        if not table.rows:
            continue
        headers = _lab_cells(table.rows[0])
        table_mapping = lab_mapping_for(headers)
        if not table_mapping:
            continue
        mapping = table_mapping
        data_rows = []
        for row_index, row in enumerate(table.rows, start=1):
            if row_index == 1:
                continue
            data_rows.append((row_index, _lab_cells(row)))
        records.extend(_records_from_header_rows(f"טבלה {table_index}", 1, headers, data_rows, "lab"))
    if not mapping:
        return None
    return {"kind": "lab", "mapping": mapping, "records": records, "sheets": []}


def _sheet_rows(worksheet) -> list[tuple[int, list[str]]]:
    found = []
    for index, row in enumerate(worksheet.iter_rows(max_col=20, values_only=True), start=1):
        found.append((index, [cell_text(value) for value in row]))
        if index >= 5000:
            break
    return found


def parse_lab_xlsx(path: Path) -> dict | None:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    records: list[dict] = []
    mapping = None
    try:
        for sheet_name in workbook.sheetnames:
            rows = _sheet_rows(workbook[sheet_name])
            header_at = None
            headers: list[str] = []
            for index, (row_number, cells) in enumerate(rows[:25]):
                maybe = lab_mapping_for(cells)
                if maybe:
                    header_at = index
                    headers = cells
                    mapping = maybe
                    break
            if header_at is None:
                continue
            data_rows = [(row_number, cells) for row_number, cells in rows[header_at + 1 :]]
            records.extend(_records_from_header_rows(sheet_name, rows[header_at][0], headers, data_rows, "lab"))
    finally:
        workbook.close()
    if not mapping:
        return None
    return {"kind": "lab", "mapping": mapping, "records": records, "sheets": []}


def parse_lab_csv(path: Path) -> dict | None:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    header_at = None
    headers: list[str] = []
    mapping = None
    for index, cells in enumerate(rows[:25]):
        maybe = lab_mapping_for(cells)
        if maybe:
            header_at = index
            headers = cells
            mapping = maybe
            break
    if not mapping or header_at is None:
        return None
    data_rows = [(index + 1, cells) for index, cells in enumerate(rows[header_at + 1 :], start=header_at + 1)]
    return {"kind": "lab", "mapping": mapping, "records": _records_from_header_rows("csv", header_at + 1, headers, data_rows, "lab"), "sheets": ["csv"]}


def _plant_headers(raw: dict) -> bool:
    keys = {header_key(key) for key in raw}
    return "item" in keys and "description" in keys


def parse_plant_upload(path: Path) -> dict | None:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        records, sheets = read_plant_xlsx(path)
    elif suffix == ".csv":
        records, sheets = read_plant_csv(path)
    else:
        return None
    if not records or not _plant_headers(records[0]["raw"]):
        return None
    return {"kind": "plant", "mapping": PLANT_MAPPING, "records": records, "sheets": sheets}


def parse_uploaded_inventory(path: Path) -> dict | None:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_lab_docx(path)
    if suffix == ".xlsx":
        return parse_lab_xlsx(path) or parse_plant_upload(path)
    if suffix == ".csv":
        return parse_lab_csv(path) or parse_plant_upload(path)
    return None


def text_looks_like_inventory(text: str) -> bool:
    haystack = (text or "").casefold()
    lab = "instrument tag" in haystack and "manufacturer" in haystack and "model" in haystack
    plant = "item" in haystack and "description" in haystack and ("remarks" in haystack or "system" in haystack)
    return lab or plant
