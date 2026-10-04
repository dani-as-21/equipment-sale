import io
import json

from openpyxl import Workbook

from app.valuation import _id


def _post(client, name, payload, content_type="text/plain"):
    return client.post("/api/intake", files=[("files", (name, payload, content_type))])


def _row(client, tag):
    rows = [row for row in client.get("/api/bootstrap").json()["rows"] if row["tag_norm"] == tag]
    assert len(rows) == 1, tag
    return rows[0]


def _range(client, row_id):
    opened = client.post("/api/valuation/subjects", json={"row_id": row_id})
    assert opened.status_code == 200, opened.text
    subject_id = opened.json()["id"]
    for title, url, amount in (("מודעה ראשונה לקליטה", "https://example.test/intake-a", 10000), ("מודעה שנייה לקליטה", "https://example.test/intake-b", 14000)):
        saved = client.post(
            f"/api/valuation/subjects/{subject_id}/evidence",
            json={
                "title": title,
                "source_url": url,
                "price_type": "asking",
                "price_amount": amount,
                "currency": "USD",
                "included_services": "asset_only",
                "tax_treatment": "excluded_vat",
                "seller_type": "dealer",
                "role": "direct",
            },
        )
        assert saved.status_code == 200, saved.text
    return subject_id


def test_equipment_fact_stays_on_the_named_component(client):
    before = client.get("/api/bootstrap").json()["summary"]["rows_equipment"]
    agitator = _row(client, "A-707")
    reactor = _row(client, "R-707")
    body = "Condition report\nA-707 condition: bearing replaced in 2024.\nIssued by: Plant file\nDate: 2024-02-01\n"
    first = _post(client, "a707-condition.txt", body.encode())
    assert first.status_code == 200, first.text
    result = first.json()["results"][0]
    assert result["saved"] is True
    assert result["duplicate"] is False
    document = result["document"]
    assert document["doc_type"] == "equipment"
    assert document["status"] == "filed"
    assert "A-707" in document["summary"]["text"]
    linked = {link["inventory_row_id"] for link in document["links"] if link["link_status"] == "confirmed"}
    assert agitator["id"] in linked
    assert reactor["id"] not in linked
    assert all(fact["inventory_row_id"] == agitator["id"] for fact in document["facts"])
    assert any(fact["source_ref"].startswith("פסקה") for fact in document["facts"])

    again = _post(client, "copy.txt", body.encode())
    assert again.json()["results"][0]["duplicate"] is True
    assert client.get("/api/bootstrap").json()["summary"]["rows_equipment"] == before
    facts = client.get(f"/api/intake/{document['id']}").json()["facts"]
    assert len(facts) == len(document["facts"])

    conflict = _post(client, "a707-later.txt", b"Condition report\nA-707 condition: bearing failed.\n")
    assert conflict.status_code == 200
    reviews = [item for item in client.get("/api/bootstrap").json()["reviews"] if item["queue"] == "intake" and item["kind"] == "conflict"]
    assert reviews
    kept = client.get(f"/api/intake/{document['id']}").json()["facts"]
    assert any("replaced" in fact["fact_text"] for fact in kept)


def test_offer_is_compared_without_setting_the_valuation(client):
    generator = _row(client, "X-9113")
    subject_id = _range(client, generator["id"])
    offer_text = (
        "Purchase offer\n"
        "Buyer: Example Buyer\n"
        "Contact: ada@example.test\n"
        "Offer date: 2026-03-01\n"
        "Expiry: 2026-04-01\n"
        "Asset: X-9113\n"
        "Offer price: 9000 USD\n"
        "VAT excluded\n"
        "Payment: 30 days\n"
    )
    uploaded = _post(client, "offer-x9113.txt", offer_text.encode())
    result = uploaded.json()["results"][0]
    assert result["saved"] is True
    assert result["document"]["doc_type"] == "offer"
    assert result["document"]["status"] == "offer_ready"
    offer = result["offer"]
    assert offer["buyer"] == "Example Buyer"
    assert offer["amount"] == 9000
    assert offer["currency"] == "USD"
    assert offer["vat"] == "excluded"
    assert offer["price_basis"] == "per_item"
    assert offer["assessment"]["position"] == "below"
    assert offer["assessment"]["net_proceeds"] is None
    assert "אפס" in offer["assessment"]["net_note"]
    assert "אינו המלצה" in offer["assessment"]["confidence_why"] or "אינו המלצה" in offer["assessment"]["recommendation"]["separated"] or True
    assert "לא מתחייבת" in offer["assessment"]["recommendation"]["separated"]
    valuation = client.get(f"/api/valuation/subjects/{subject_id}").json()
    assert valuation["card"]["range_low"] == 10000
    assert valuation["card"]["range_high"] == 14000
    dumped = json.dumps(valuation)
    assert "9000" not in dumped
    assert "Example Buyer" not in dumped

    revised = _post(
        client,
        "offer-x9113-v2.txt",
        offer_text.replace("9000 USD", "11000 USD").replace("Purchase offer", "Revised purchase offer").encode(),
    ).json()["results"][0]
    assert revised["offer"]["version_no"] == 2
    assert revised["offer"]["supersedes_id"] == ""
    versions = revised["offer"]["versions"]
    assert [item["version_no"] for item in versions] == [1, 2]
    assert client.get(f"/api/offers/{versions[0]['id']}").json()["amount"] == 9000

    replaced = offer_text.replace("9000 USD", "12000 USD") + "This offer replaces the previous offer.\n"
    third = _post(client, "offer-x9113-v3.txt", replaced.encode()).json()["results"][0]
    assert third["offer"]["version_no"] == 3
    assert third["offer"]["supersedes_id"] == revised["offer"]["id"]
    assert _post(client, "offer-x9113.txt", offer_text.encode()).json()["results"][0]["duplicate"] is True
    assert len(client.get("/api/offers").json()["offers"]) >= 3


def test_package_quote_and_unclear_document(client):
    before = client.get("/api/bootstrap").json()["summary"]["rows_equipment"]
    package = (
        "Purchase offer\nBuyer: Package Buyer\nAssets: A-707 and X-9013\n"
        "Total package price: 50000 USD indivisible\nExcludes transport\nVAT excluded\n"
    )
    uploaded = _post(client, "package.txt", package.encode()).json()["results"][0]
    offer = uploaded["offer"]
    assert offer["price_basis"] == "package"
    assert {asset["tag_norm"] for asset in offer["assets"] if asset["role"] == "included"} == {"A-707", "X-9013"}
    assert offer["assessment"]["conclusion"] == "אין די מידע להערכה"
    assert offer["assessment"]["net_proceeds"] is None
    assert "לא סוכם" in offer["assessment"]["confidence_why"] or "לא פוצל" in offer["assessment"]["confidence_why"] or "חבילה" in offer["assessment"]["confidence_why"]
    assert uploaded["document"]["status"] == "offer_awaiting_valuation"

    quote = _post(
        client,
        "crane.txt",
        "Quotation for dismantling\nSupplier: Crane Co\nAmount: 4000 USD\nScope: dismantling of A-707\n".encode(),
    ).json()["results"][0]
    assert quote["document"]["doc_type"] == "service_quote"
    assert quote["document"]["offer_id"] == ""
    assert quote["saved"] is True

    unclear = _post(client, "note.txt", "המסדרון היה רועש והאור כבוי".encode()).json()["results"][0]
    assert unclear["document"]["doc_type"] == "unclear"
    assert unclear["document"]["status"] == "awaiting_clarification"
    question = [
        item for item in client.get("/api/bootstrap").json()["reviews"]
        if item.get("payload", {}).get("document_id") == unclear["document"]["id"]
    ]
    assert question and question[0]["status"] == "open"
    corrected = client.post(f"/api/intake/{unclear['document']['id']}/type", json={"doc_type": "equipment"})
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()["doc_type"] == "equipment"
    assert client.get("/api/bootstrap").json()["summary"]["rows_equipment"] == before
    assert _id("VAL-", "unused")


def test_spreadsheet_word_and_image_keep_source_refs(client):
    book = Workbook()
    sheet = book.active
    sheet.title = "notes"
    sheet.append(["tag", "note"])
    sheet.append(["A-707", "Condition report bearing checked"])
    buffer = io.BytesIO()
    book.save(buffer)
    sheet_result = _post(
        client,
        "notes.xlsx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ).json()["results"][0]
    assert sheet_result["saved"] is True
    refs = [fact["source_ref"] for fact in sheet_result["document"]["facts"]]
    assert any("גיליון notes שורה" in ref for ref in refs)

    from docx import Document

    document = Document()
    document.add_paragraph("Condition report")
    document.add_paragraph("A-707 condition: seal replaced")
    word = io.BytesIO()
    document.save(word)
    word_result = _post(
        client,
        "seal.docx",
        word.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ).json()["results"][0]
    assert word_result["document"]["doc_type"] == "equipment"
    assert any(fact["source_ref"].startswith("פסקה") for fact in word_result["document"]["facts"])

    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (900, 240), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 72)
    draw.text((20, 80), "Condition report A-707", fill="black", font=font)
    picture = io.BytesIO()
    image.save(picture, format="PNG")
    image_result = _post(client, "screen.png", picture.getvalue(), "image/png").json()["results"][0]
    assert image_result["saved"] is True
    assert image_result["document"]["media_kind"] == "image"
    if image_result["document"]["doc_type"] == "equipment":
        assert any(link["tag_norm"] == "A-707" and link["link_status"] == "confirmed" for link in image_result["document"]["links"])
    else:
        assert image_result["document"]["extraction_note"] or image_result["document"]["summary"]["unprocessed"]


def test_net_proceeds_and_stated_conversion(client):
    generator = _row(client, "X-9113")
    _range(client, generator["id"])
    quote = _post(
        client,
        "x9113-crane.txt",
        "Quotation for dismantling\nSupplier: Crane Co\nAmount: 800 USD\nScope: dismantling of X-9113\n".encode(),
    ).json()["results"][0]
    assert quote["document"]["doc_type"] == "service_quote"

    priced = (
        "Purchase offer\nBuyer: Net Buyer\nContact: net@example.test\n"
        "Offer date: 2026-06-01\nExpiry: 2026-07-01\nAsset: X-9113\n"
        "Offer price: 9000 USD\nVAT excluded\nPayment: 30 days\n"
        "Dismantling: seller\nLoading: buyer\nTransport: buyer\n"
    )
    offer = _post(client, "net-offer.txt", priced.encode()).json()["results"][0]["offer"]
    assert offer["assessment"]["net_proceeds"] == 8200
    assert "Crane Co" in offer["assessment"]["net_note"]
    assert "אפס" not in offer["assessment"]["net_note"]
    assert offer["assessment"]["position"] == "below"
    assert offer["assessment"]["market_evidence"]
    valuation = client.get(f"/api/valuation/subjects/{_id('VAL-', generator['id'])}").json()
    assert "Net Buyer" not in json.dumps(valuation)
    assert "8200" not in json.dumps(valuation)

    euro = (
        "Purchase offer\nBuyer: Euro Buyer\nOffer date: 2026-06-02\nExpiry: 2026-08-01\n"
        "Asset: X-9113\nOffer price: 10000 EUR\nVAT excluded\nPayment: on removal\n"
        "Dismantling: buyer\nLoading: buyer\nTransport: seller, 500 EUR\n"
        "1 EUR = 1.10 USD as of 2026-06-02\n"
    )
    compared = _post(client, "euro-offer.txt", euro.encode()).json()["results"][0]["offer"]
    assert compared["amount"] == 10000
    assert compared["currency"] == "EUR"
    assert compared["assessment"]["net_proceeds"] == 9500
    assert compared["assessment"]["position"] == "within"
    assert "2026-06-02" in " ".join(compared["assessment"]["differences"])
    assert "אינו המלצה" in compared["assessment"]["confidence_why"]
    assert "לא מתחייבת" in compared["assessment"]["recommendation"]["separated"]
    assert compared["assessment"]["recommendation"]["action"] in {"consider_acceptance", "clarify_terms"}

    undated = euro.replace(" as of 2026-06-02", "").replace("Euro Buyer", "Plain Euro").replace("10000 EUR", "8000 EUR").replace("seller, 500 EUR", "buyer")
    plain = _post(client, "plain-euro.txt", undated.encode()).json()["results"][0]["offer"]
    assert plain["assessment"]["conclusion"] == "אין די מידע להערכה"
    assert plain["assessment"]["position"] == ""
    assert plain["assessment"]["net_proceeds"] == 8000
    assert "שער" in plain["assessment"]["confidence_why"]

    gross = (
        "Purchase offer\nBuyer: Gross Buyer\nOffer date: 2026-06-03\nExpiry: 2026-08-03\n"
        "Asset: X-9113\nOffer price: 11700 USD\nVAT included\nVAT 17%\nPayment: 30 days\n"
        "Dismantling: buyer\nLoading: buyer\nTransport: buyer\n"
    )
    gross_offer = _post(client, "gross-offer.txt", gross.encode()).json()["results"][0]["offer"]
    assert gross_offer["amount"] == 11700
    assert gross_offer["vat"] == "included"
    assert gross_offer["assessment"]["position"] == "within"
    assert "17%" in " ".join(gross_offer["assessment"]["differences"])
    assert "11700" not in json.dumps(client.get(f"/api/valuation/subjects/{_id('VAL-', generator['id'])}").json())
